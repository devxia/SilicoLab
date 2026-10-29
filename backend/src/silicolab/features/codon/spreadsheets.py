"""Local-only spreadsheet preview and synthesis-workbook export."""

import csv
from copy import copy
from pathlib import Path
from datetime import datetime
import io
from zipfile import ZipFile

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_SHEETS = 20
MAX_CELLS = 100000


class SpreadsheetError(ValueError):
    pass


def inspect_table(content: bytes, filename: str) -> dict:
    """Read cached values only. Never execute formulas, macros, or external links."""
    if len(content) > MAX_FILE_BYTES:
        raise SpreadsheetError("文件超过 5 MB，请分批导入")
    if filename.startswith("~$"):
        raise SpreadsheetError("这是 Excel 临时锁文件，请选择原始工作簿")
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension in {"csv", "tsv"}:
        return inspect_delimited(content, filename, extension)
    if extension != "xlsx":
        raise SpreadsheetError("支持 .xlsx、.csv 和 .tsv；旧版 .xls 请先另存为 .xlsx")
    from openpyxl import load_workbook
    workbook = None
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise SpreadsheetError("工作簿解压后过大，请删除无关工作表或拆分文件")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > MAX_SHEETS:
            raise SpreadsheetError(f"工作簿超过 {MAX_SHEETS} 个工作表，请先拆分")
        sheets = []
        cells = 0
        for sheet in workbook.worksheets:
            # Ignore stale, inflated worksheet dimensions; count actual cells instead.
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [cell_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if len(values) > MAX_COLUMNS:
                    raise SpreadsheetError(f"工作表「{sheet.title}」超过 {MAX_COLUMNS} 列，请删除无关列")
                rows.append(values)
                cells += max(1, len(values))
                if len(rows) > MAX_ROWS or cells > MAX_CELLS:
                    raise SpreadsheetError(f"工作簿过大：每页最多 {MAX_ROWS} 行、总计 {MAX_CELLS:,} 个单元格，请先拆分")
            while rows and not any(rows[-1]):
                rows.pop()
            sheets.append({"name": sheet.title, "rows": rows})
        if not any(sheet["rows"] for sheet in sheets):
            raise SpreadsheetError("工作簿中没有可读取的数据")
        return {"filename": filename, "sheets": sheets}
    except SpreadsheetError:
        raise
    except Exception as exc:
        raise SpreadsheetError("无法读取工作簿，请确认文件是未加密的有效 .xlsx 文件") from exc
    finally:
        if workbook is not None:
            workbook.close()


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value)


def inspect_delimited(content, filename, extension):
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SpreadsheetError("无法识别文本编码，请保存为 UTF-8 CSV 或 TSV")
    if extension == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
    rows = []
    try:
        for row in csv.reader(io.StringIO(text), delimiter=delimiter):
            rows.append(row)
            if len(row) > MAX_COLUMNS or len(rows) > MAX_ROWS or sum(map(len, rows)) > MAX_CELLS:
                raise SpreadsheetError("表格过大，请先拆分为最多 2,000 行、100 列的文件")
    except csv.Error as exc:
        raise SpreadsheetError("无法解析文本表格，请检查分隔符与引号") from exc
    if not any(any(cell.strip() for cell in row) for row in rows):
        raise SpreadsheetError("表格中没有数据")
    return {"filename": filename, "sheets": [{"name": "数据", "rows": rows}]}


def write_text(sheet, row, column, value):
    """Force untrusted names/proteins to Excel text, including names beginning '='."""
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def synthesis_workbook(records: list[dict]) -> bytes:
    """Fill the blank, data-free copy of the user's synthesis-sheet format."""
    from openpyxl import load_workbook
    template = Path(__file__).resolve().parent / "resources" / "synthesis.xlsx"
    workbook = load_workbook(template)
    synthesis = workbook.active
    row_styles = [copy(synthesis.cell(2, column)._style) for column in range(1, 4)]
    synthesis.delete_rows(2, synthesis.max_row - 1)
    valid = [record for record in records if record["status"] == "completed" and record.get("dna")]
    for index, record in enumerate(valid, 1):
        values = [index, record["name"], record["dna"]]
        for column, value in enumerate(values, 1):
            cell = write_text(synthesis, index + 1, column, value)
            cell._style = copy(row_styles[column - 1])
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                cell.comment = None
    content = io.BytesIO()
    workbook.save(content)
    workbook.close()
    return content.getvalue()
