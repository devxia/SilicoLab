import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from silicolab.features.codon import client as codon
from silicolab.features.codon.engines import engine_metadata
from app_fixture import create_app


class EngineIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(Path(self.temp.name) / 'db.sqlite3', self.remote, delay=0)
        self.calls = []
        self.client = TestClient(self.app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.temp.cleanup()

    def remote(self, payload):
        self.calls.append(payload)
        return {'success': '1', 'seq': ''.join({'M': 'ATG', 'K': 'AAG', '*': 'TAA'}[aa] for aa in payload['sequence']), 'log_id': 'remote-test'}

    def submit(self, **updates):
        return self.client.post('/api/jobs', json={
            'name': 'engine test', 'engine': 'local', 'species': 'D2P(default)',
            'avoid': '', 'records': [{'name': 'protein', 'sequence': 'MK*'}], **updates,
        })

    def wait(self, job_id):
        for _ in range(300):
            job = self.client.get('/api/jobs/' + job_id).json()
            if job['status'] not in {'queued', 'running', 'stopping'}:
                return job
            time.sleep(0.01)
        self.fail('Job did not complete')

    def test_local_job_is_fully_offline_and_persists_engine_fingerprints(self):
        with patch('silicolab.features.codon.client.request_json', side_effect=AssertionError('network forbidden')) as network:
            response = self.submit(avoid='AAG')
            self.assertEqual(response.status_code, 201)
            job = self.wait(response.json()['id'])
            self.assertEqual(job['status'], 'completed')
            self.assertEqual(job['engine'], 'local')
            info = job['engine_info']
            self.assertEqual(info['version'], 'd2p-local-experimental-1')
            self.assertFalse(info['network_used'])
            self.assertEqual(len(info['mapping_sha256']), 64)
            self.assertEqual(len(info['code_sha256']), 64)
            self.assertEqual(info['constraint_strategy'], 'min_changed_codons_then_nucleotide_edits_v1')
            record = job['records'][0]
            self.assertEqual(record['dna'], 'ATGAAATAA')
            self.assertTrue(record['checks']['translation'])
            self.assertFalse(record['response']['network_used'])
            self.assertEqual(record['attempts'][0]['engine'], 'local')
            self.assertEqual(record['attempts'][0]['engine_info'], info)
            network.assert_not_called()
        self.assertEqual(self.calls, [])

    def test_remote_remains_default_when_engine_omitted(self):
        response = self.client.post('/api/jobs', json={'name': 'old request', 'records': [{'name': 'a', 'sequence': 'MK'}]})
        job = self.wait(response.json()['id'])
        self.assertEqual(job['engine'], 'remote')
        self.assertTrue(job['engine_info']['network_used'])
        self.assertEqual(len(self.calls), 1)
        self.assertNotIn('engine', self.calls[0])
        self.assertEqual(job['records'][0]['response']['log_id'], 'remote-test')

    def test_non_d2p_local_and_unknown_engine_rejected_before_job_creation(self):
        for changes in [{'species': 'Escherichia_coli'}, {'engine': 'unknown'}]:
            self.assertEqual(self.submit(**changes).status_code, 422)
        self.assertEqual(self.calls, [])
        self.assertEqual(self.client.get('/api/jobs').json(), [])

    def test_local_motif_size_limit_is_validated_before_submission(self):
        response = self.submit(avoid='A' * 257)
        self.assertEqual(response.status_code, 422)
        self.assertIn('256', response.text)
        self.assertEqual(self.calls, [])

    def test_local_infeasibility_and_retry_never_fall_back_to_remote(self):
        with patch('silicolab.features.codon.client.request_json', side_effect=AssertionError('network forbidden')) as network:
            response = self.submit(avoid='ATG')
            job = self.wait(response.json()['id'])
            self.assertEqual(job['status'], 'failed')
            self.assertIn('无同义编码', job['records'][0]['error'])
            self.assertEqual(self.client.post(f"/api/jobs/{job['id']}/retry").status_code, 200)
            retried = self.wait(job['id'])
            self.assertEqual(retried['status'], 'failed')
            self.assertEqual(len(retried['records'][0]['attempts']), 2)
            self.assertTrue(all(attempt['engine'] == 'local' for attempt in retried['records'][0]['attempts']))
            network.assert_not_called()
        self.assertEqual(self.calls, [])

    def test_presets_retain_engine_and_reject_incompatible_local_host(self):
        response = self.client.post('/api/presets', json={'name': 'local preset', 'engine': 'local', 'species': 'D2P(default)', 'avoid': 'GAATTC'})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(self.client.get('/api/presets').json()[0]['engine'], 'local')
        invalid = self.client.post('/api/presets', json={'name': 'bad', 'engine': 'local', 'species': 'Escherichia_coli'})
        self.assertEqual(invalid.status_code, 422)

    def test_legacy_jobs_and_presets_are_explicitly_remote_on_read(self):
        response = self.submit(engine='remote')
        job = self.wait(response.json()['id'])
        job.pop('engine')
        job.pop('engine_info')
        self.app.state.store.save(job)
        loaded = self.client.get('/api/jobs/' + job['id']).json()
        self.assertEqual(loaded['engine'], 'remote')
        self.assertTrue(loaded['engine_info']['legacy'])
        self.assertEqual(self.client.get('/api/jobs').json()[0]['engine'], 'remote')
        self.app.state.store.save_preset({'id': 'old', 'name': 'old preset', 'species': 'D2P(default)', 'avoid': ''})
        self.assertEqual(self.client.get('/api/presets').json()[0]['engine'], 'remote')

    def test_paused_remote_does_not_block_new_local_jobs(self):
        response = self.submit(engine='remote')
        job = self.wait(response.json()['id'])
        self.app.state.store.mutate(job['id'], lambda item: item.update(status='paused'))
        local_response = self.submit()
        self.assertEqual(local_response.status_code, 201)
        self.assertEqual(self.wait(local_response.json()['id'])['status'], 'completed')
        self.assertEqual(self.submit(engine='remote').status_code, 409)
        self.assertEqual(len(self.calls), 1)

    def test_remote_rate_limit_pauses_only_remote_queue(self):
        entered, release = threading.Event(), threading.Event()
        def limited(payload):
            entered.set()
            release.wait(3)
            try:
                raise HTTPError('https://example.test', 429, 'limit', {}, None)
            except HTTPError as exc:
                raise codon.OptimizationError('HTTP 429') from exc
        self.app.state.worker.optimizer = limited
        response = self.submit(engine='remote')
        self.assertTrue(entered.wait(2))
        local_response = self.submit()
        remote_response = self.submit(engine='remote')
        release.set()
        self.assertEqual(self.wait(response.json()['id'])['status'], 'paused')
        self.assertEqual(self.wait(local_response.json()['id'])['status'], 'completed')
        self.assertEqual(self.wait(remote_response.json()['id'])['status'], 'paused')

    def test_retry_rejects_version_drift_without_modifying_successful_results(self):
        response = self.submit(avoid='ATG')
        job = self.wait(response.json()['id'])
        self.app.state.store.mutate(job['id'], lambda item: item['engine_info'].update(version='older-version'))
        response = self.client.post(f"/api/jobs/{job['id']}/retry")
        self.assertEqual(response.status_code, 409)
        self.assertIn('版本已变更', response.text)
        self.assertEqual(self.calls, [])
        self.assertEqual(self.client.get('/api/jobs/' + job['id']).json()['status'], 'failed')

    def test_local_exports_keep_clean_xlsx_and_include_engine_in_zip(self):
        response = self.submit(stop_policy='add', records=[{'name': 'added', 'sequence': 'MK'}])
        job = self.wait(response.json()['id'])
        exported = self.client.get(f"/api/jobs/{job['id']}/export?format=xlsx")
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book.sheetnames, ['基因合成'])
        self.assertEqual(book.active['C2'].value, 'ATGAAGTAA')
        self.assertFalse(any(cell.comment for sheet in book for row in sheet for cell in row))
        book.close()
        bundle = self.client.get(f"/api/jobs/{job['id']}/export?format=zip")
        with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
            saved = json.loads(archive.read('record.json'))
            self.assertEqual(saved['engine_info'], engine_metadata('local'))
            self.assertEqual(saved['records'][0]['attempts'][0]['engine'], 'local')
        self.assertEqual(self.calls, [])

    def test_cached_species_request_is_offline_even_when_cache_is_empty(self):
        with patch('silicolab.server.codon.list_species', side_effect=AssertionError('No network')) as remote:
            response = self.client.get('/api/species?cached=true')
            self.assertEqual(response.status_code, 200)
            self.assertIn('D2P(default)', response.json()['items'])
            self.assertIsNone(response.json()['online'])
            remote.assert_not_called()


if __name__ == '__main__':
    unittest.main()
