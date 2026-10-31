"""Actual generated Python through kernel + Seatbelt + app supervisor with synthetic inputs."""
import base64
import io
import json
import os
from pathlib import Path
import signal
import socket
import tempfile
import time
import unittest
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app_fixture import create_app
from test_model_configuration import FakeModelEndpoint, MemorySecrets
from test_workbench_science import wait_task
from test_workbench_recovery import wait_state

RUNTIME_APP = Path(__file__).resolve().parents[2] / 'release/PythonBoundaryProbe.app'
SCRIPT = '''import csv,json,statistics,math
from pathlib import Path
from openpyxl import Workbook
values=[float(r['value']) for r in csv.DictReader((INPUTS/'input-001.csv').open())]
result={'count':len(values),'sum':math.fsum(values),'mean':statistics.mean(values)}
Path('summary.json').write_text(json.dumps(result))
book=Workbook();sheet=book.active;sheet.append(['count','sum','mean']);sheet.append([result['count'],result['sum'],result['mean']]);book.save('summary.xlsx')
'''


@unittest.skipUnless(RUNTIME_APP.exists(), 'Explicitly build tools/python-runtime before restricted-analysis integration checks')
class RestrictedAnalysisTests(unittest.TestCase):
    def setup(self, client, root, model=None, csv_content=b'value\n1\n2\n6\n'):
        project = root / 'project'; project.mkdir()
        opened = client.post('/api/workbench/projects', json={'path': str(project)}).json()
        url = '/api/workbench/projects/' + opened['id']
        chat = client.post(url + '/conversations', json={'title': 'temporary numeric analysis'}).json()['id']
        source = client.post(url + '/files', json={'name': 'artificial.csv', 'contentBase64': base64.b64encode(csv_content).decode()}).json()
        fixed = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [source['id']]}).json()
        if model:
            config = client.post('/api/workbench/model-configs', json={'name': 'artificial analysis model', 'protocol': model.protocol, 'baseUrl': model.url, 'modelId': 'artificial-model', 'apiKey': model.key}).json()
            client.put(url + f'/conversations/{chat}/model', json={'configurationId': config['id']})
        body = {'conversationId': chat, 'requestId': 'analysis', 'goal': '临时分析人工CSV数值摘要，用预装库生成JSON和Excel', 'workflow': 'python_analysis', 'authorizedTools': ['python_analysis'], 'inputVersionId': fixed['id']}
        response = client.post(url + '/tasks', json=body)
        self.assertEqual(response.status_code, 201, response.text)
        return url, chat, response.json()['id'], fixed, source, opened

    def test_model_generated_script_runs_and_exports_independently_verified_summary(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3', delay=0)) as client:
                url,chat,identity,fixed,source,_=self.setup(client,root)
                task=wait_task(client,url,chat,identity)
                execution=task['toolExecutions'][0]
                self.assertEqual(execution['tool'],'python_analysis')
                self.assertEqual(execution['status'],'completed',execution)
                result=execution['result']
                self.assertTrue(result['checks']['independent_numeric_summary'])
                self.assertEqual(result['summary'],{'count':3,'sum':9.0,'mean':3.0})
                self.assertEqual(result['inputVersionId'],fixed['id'])
                self.assertEqual(result['environment']['python'],'3.13.18')
                artifacts={a['name']:a for a in execution['artifacts']}
                self.assertIn('analysis.py',artifacts)
                self.assertIn('summary.json',artifacts)
                self.assertIn('summary.xlsx',artifacts)
                base=url+f'/tasks/{identity}/artifacts/{execution["id"]}'
                workbook=load_workbook(io.BytesIO(client.get(base+'/summary.xlsx').content))
                self.assertEqual(list(workbook.active.values),[('count','sum','mean'),(3,9,3)])
                workbook.close()
                self.assertIn('statistics',client.get(base+'/analysis.py').text)
                self.assertEqual((root/'project'/source['path']).read_bytes(),b'value\n1\n2\n6\n')
                self.assertEqual(client.get('/api/jobs').json(),[])

    def test_model_receives_verified_numeric_summary_and_artifact_references_on_fresh_and_reused_tool(self):
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-summary',mode=['normal','duplicate_once','normal'],tool_name='python_analysis',tool_arguments={'code':SCRIPT}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                task=wait_task(client,url,chat,identity)
                self.assertEqual(len(task['toolExecutions']),1,'same completed script must be reused')
                expected={a['name']:a for a in task['toolExecutions'][0]['artifacts']}
                self.assertEqual(len(model.requests),3)
                for request in model.requests[1:]:
                    tool_results=[json.loads(item['output']) for item in request['body']['input'] if item.get('type')=='function_call_output']
                    for result in tool_results:
                        self.assertEqual(result.get('summary'),{'count':3,'sum':9,'mean':3})
                        self.assertTrue(result.get('independentlyVerified'))
                        refs={a['name']:a for a in result['artifacts']}
                        self.assertEqual(set(refs),{'analysis.py','summary.json','summary.xlsx'})
                        for name,ref in refs.items():
                            self.assertEqual(ref['sha256'],expected[name]['sha256']);self.assertNotIn('path',ref)
                        self.assertNotIn(SCRIPT,json.dumps(result))

    def test_supervisor_rejects_small_workbook_with_oversized_dimensions_before_unbounded_parse(self):
        # Each adversarial file is tiny; preflight must reject coordinates before parsing rows.
        for dimension,row in [('A1:XFD3',3),('A1:C100',100),('A1:XFD10000',10000),('A1:C1000000000',1000000000),('A1:C2',1000000000)]:
            code=SCRIPT+f'''
import re,zipfile
path=Path('summary.xlsx')
with zipfile.ZipFile(path) as source: entries={{n:source.read(n) for n in source.namelist()}}
xml=entries['xl/worksheets/sheet1.xml'].decode()
xml=re.sub(r'<dimension[^>]*/>', '<dimension ref="{dimension}"/>', xml)
xml=xml.replace('</sheetData>', '<row r="{row}"/></sheetData>')
entries['xl/worksheets/sheet1.xml']=xml.encode()
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as output:
 for name,data in entries.items():output.writestr(name,data)
'''
            with self.subTest(dimension=dimension),tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-dimensions',tool_name='python_analysis',tool_arguments={'code':code}) as model:
                root=Path(directory)
                with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                    url,chat,identity,_,_,_=self.setup(client,root,model)
                    done=wait_task(client,url,chat,identity)
                    execution=done['toolExecutions'][0]
                    self.assertEqual(execution['status'],'failed')
                    self.assertIn('行列范围',execution['result']['error'])
                    self.assertEqual(client.post(url+f'/tasks/{identity}/pause').status_code,200)

    def test_relationship_cannot_redirect_to_unchecked_binary_named_worksheet(self):
        code=SCRIPT+'''
import zipfile,re
path=Path('summary.xlsx')
with zipfile.ZipFile(path) as source: entries={n:source.read(n) for n in source.namelist()}
real=entries['xl/worksheets/sheet1.xml'].decode()
real=real.replace('<c r="A2" t="n"><v>3</v></c>', '<c r="A2" t="n"><f>1+2</f><v>3</v></c>')
entries['xl/actual.bin']=real.encode()
entries['xl/_rels/workbook.xml.rels']=entries['xl/_rels/workbook.xml.rels'].replace(b'/xl/worksheets/sheet1.xml',b'/xl/actual.bin')
entries['[Content_Types].xml']=entries['[Content_Types].xml'].replace(b'/xl/worksheets/sheet1.xml',b'/xl/actual.bin')
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as output:
 for name,data in entries.items(): output.writestr(name,data)
'''
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-relationship',tool_name='python_analysis',tool_arguments={'code':code}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                done=wait_task(client,url,chat,identity)
                execution=done['toolExecutions'][0]
                self.assertEqual(execution['status'],'failed')
                self.assertRegex(execution['result']['error'],'部件|关系|Target')

    def test_namespace_alias_cannot_reinterpret_styles_as_actual_worksheet(self):
        code=SCRIPT+'''
import zipfile,xml.etree.ElementTree as ET
path=Path('summary.xlsx')
with zipfile.ZipFile(path) as source: entries={n:source.read(n) for n in source.namelist()}
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
rns='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
relations=ET.fromstring(entries['xl/_rels/workbook.xml.rels'])
style_id=next(node.get('Id') for node in relations if node.get('Type').endswith('/styles'))
workbook=ET.fromstring(entries['xl/workbook.xml'])
alias=ET.SubElement(workbook,'sheets',{'xmlns':''})
ET.SubElement(alias,'sheet',{'name':'Alias','sheetId':'2','{'+rns+'}id':style_id})
entries['xl/workbook.xml']=ET.tostring(workbook)
styles=ET.fromstring(entries['xl/styles.xml'])
worksheet=ET.fromstring(entries['xl/worksheets/sheet1.xml'])
for node in worksheet:
 if node.tag.rsplit('}',1)[-1] in ('dimension','sheetData'): styles.append(node)
cell=styles.find('.//{'+ns+'}c[@r="A2"]')
formula=ET.Element('{'+ns+'}f');formula.text='1+2';cell.insert(0,formula)
entries['xl/styles.xml']=ET.tostring(styles)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as output:
 for name,data in entries.items(): output.writestr(name,data)
'''
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-namespace',tool_name='python_analysis',tool_arguments={'code':code}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                done=wait_task(client,url,chat,identity)
                execution=done['toolExecutions'][0]
                self.assertEqual(execution['status'],'failed',execution)
                self.assertRegex(execution['result']['error'],'命名空间|工作表|结构|样式')

    def test_date_styles_explicit_or_implicit_zero_are_not_numeric_evidence(self):
        cases = {
            'explicit_date': SCRIPT+"\nsheet['A2'].number_format='yyyy-mm-dd'\nbook.save('summary.xlsx')\n",
            'implicit_zero': SCRIPT+'''
import zipfile,xml.etree.ElementTree as ET
path=Path('summary.xlsx')
with zipfile.ZipFile(path) as source:entries={n:source.read(n) for n in source.namelist()}
ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
styles=ET.fromstring(entries['xl/styles.xml'])
styles.find(ns+'cellXfs')[0].set('numFmtId','14')
entries['xl/styles.xml']=ET.tostring(styles)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as output:
 for name,data in entries.items():output.writestr(name,data)
''',
        }
        for case,code in cases.items():
            with self.subTest(case=case),tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-date-style',tool_name='python_analysis',tool_arguments={'code':code}) as model:
                root=Path(directory)
                with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                    url,chat,identity,_,_,_=self.setup(client,root,model)
                    execution=wait_task(client,url,chat,identity)['toolExecutions'][0]
                    self.assertEqual(execution['status'],'failed')
                    self.assertIn('样式',execution['result']['error'])
                    self.assertFalse(execution['result']['checks'].get('json_and_workbook_match',False))

    def test_default_sum_uses_version_independent_compensated_definition(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0)) as client:
                url,chat,identity,_,_,_=self.setup(client,root,csv_content=b'value\n1e16\n1\n-1e16\n')
                done=wait_task(client,url,chat,identity)
                result=done['toolExecutions'][0]['result']
                self.assertEqual(result['summary']['sum'],1.0)
                self.assertAlmostEqual(result['summary']['mean'],1/3)
                script=client.get(url+f'/tasks/{identity}/artifacts/{done["toolExecutions"][0]["id"]}/analysis.py').text
                self.assertIn('math.fsum',script)

    def test_default_mean_is_independently_validated_for_scale_cancellation(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0)) as client:
                url,chat,identity,_,_,_=self.setup(client,root,csv_content=b'value\n1e100\n1e50\n1\n-1e100\n-1e50\n')
                done=wait_task(client,url,chat,identity)
                execution=done['toolExecutions'][0]
                self.assertEqual(execution['status'],'completed',execution)
                self.assertEqual(execution['result']['summary']['count'],5)
                self.assertAlmostEqual(execution['result']['summary']['mean'],0.2)
                self.assertTrue(execution['result']['checks']['json_and_workbook_match'])

    def test_script_record_name_is_reserved_and_artifact_hashes_remain_unique(self):
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-reserved',tool_name='python_analysis',tool_arguments={'code':SCRIPT+"\nPath('record.json').write_text('{}')\n"}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                done=wait_task(client,url,chat,identity)
                execution=done['toolExecutions'][0]
                self.assertEqual(execution['status'],'failed')
                self.assertIn('保留',execution.get('error',''))
                names=[a['name'] for a in execution['artifacts']]
                self.assertEqual(len(names),len(set(names)))
                for name in names:
                    self.assertEqual(client.get(url+f'/tasks/{identity}/artifacts/{execution["id"]}/{name}').status_code,200)

    def test_full_log_size_over_limit_is_rejected_not_silently_truncated(self):
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-log-limit',tool_name='python_analysis',tool_arguments={'code':SCRIPT+"\nprint('x' * (2*1024*1024))\n"}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                done=wait_task(client,url,chat,identity)
                execution=done['toolExecutions'][0]
                self.assertEqual(execution['status'],'failed')
                self.assertIn('1MiB',execution.get('error',''))

    def test_log_output_alias_cannot_hide_retained_descriptor_limit(self):
        code=SCRIPT+"\nimport sys,os\nprint('x'*(2*1024*1024));sys.stdout.flush()\nos.unlink('stdout.txt')\nPath('stdout.txt').write_text('small replacement')\n"
        with tempfile.TemporaryDirectory() as directory,FakeModelEndpoint('openai-responses','artificial-retained-log',tool_name='python_analysis',tool_arguments={'code':code}) as model:
            root=Path(directory)
            with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                done=wait_task(client,url,chat,identity)
                self.assertEqual(done['toolExecutions'][0]['status'],'failed')
                self.assertIn('1MiB',done['toolExecutions'][0]['error'])

    def test_failed_missing_or_wrong_outputs_never_become_computed_success(self):
        for code in ('raise RuntimeError("artificial runtime failure")', 'pass', 'from pathlib import Path\nPath("summary.json").write_text(\'{"count":3,"sum":999,"mean":333}\')'):
            with self.subTest(code=code),tempfile.TemporaryDirectory() as directory, FakeModelEndpoint('openai-responses','artificial-failure',tool_name='python_analysis',tool_arguments={'code':code}) as model:
                root=Path(directory)
                with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                    url,chat,identity,_,_,_=self.setup(client,root,model)
                    done=wait_task(client,url,chat,identity)
                    execution=done['toolExecutions'][0]
                    self.assertEqual(execution['status'],'failed')
                    self.assertFalse(execution['result']['checks'].get('independent_numeric_summary',False))
                    self.assertTrue(any(a['name']=='analysis.py' for a in execution['artifacts']))

    def test_missing_verified_runtime_blocks_without_fallback(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            moved=RUNTIME_APP.with_name('PythonBoundaryProbe.artificial-hidden.app')
            RUNTIME_APP.rename(moved)
            try:
                with TestClient(create_app(root/'app.sqlite3',delay=0)) as client:
                    url,chat,identity,_,_,_=self.setup(client,root)
                    done=wait_task(client,url,chat,identity)
                    self.assertEqual(done['toolExecutions'][0]['status'],'failed')
                    self.assertIn('未预装',done['toolExecutions'][0]['error'])
                    self.assertFalse(any(a['name']=='summary.xlsx' for a in done['toolExecutions'][0]['artifacts']))
            finally:
                moved.rename(RUNTIME_APP)

    def test_adversarial_script_cannot_get_model_key_change_original_network_fork_or_install(self):
        with tempfile.TemporaryDirectory() as directory, socket.socket() as listener:
            root=Path(directory); sentinel=root/'artificial-credential.txt';sentinel.write_text('artificial-analysis-model-key')
            listener.bind(('127.0.0.1',0));listener.listen()
            code=f'''import os,socket,json,sys,subprocess
from pathlib import Path
checks={{}}
for name,action in [('outside',lambda:Path({str(sentinel)!r}).read_text()),('input-write',lambda:(INPUTS/'input-001.csv').write_text('changed')),('network',lambda:socket.create_connection(('127.0.0.1',{listener.getsockname()[1]}),timeout=.2)),('fork',lambda:os.fork()),('install',lambda:__import__('ensurepip').bootstrap())]:
 try: action();checks[name]='ALLOWED'
 except (OSError,ImportError) as e: checks[name]=getattr(e,'errno',None) or 'unavailable'
checks['environment_key']=os.environ.get('OPENAI_API_KEY')
Path('denied.json').write_text(json.dumps(checks))
'''
            with FakeModelEndpoint('openai-responses','artificial-analysis-model-key',tool_name='python_analysis',tool_arguments={'code':code}) as model, patch.dict('os.environ',{'OPENAI_API_KEY':'artificial-host-key'}):
                with TestClient(create_app(root/'app.sqlite3',delay=0,credential_store=MemorySecrets())) as client:
                    url,chat,identity,_,_,_=self.setup(client,root,model)
                    task=wait_task(client,url,chat,identity)
                    execution=task['toolExecutions'][0]
                    self.assertEqual(execution['status'],'failed')
                    artifact=next(a for a in execution['artifacts'] if a['name']=='denied.json')
                    checks=json.loads((root/'project'/artifact['path']).read_text())
                    for name in ('outside','input-write','network','fork','install'):
                        self.assertIn(checks[name],(1,13,'unavailable'),name)
                    self.assertIsNone(checks['environment_key'])
                    self.assertEqual(sentinel.read_text(),'artificial-analysis-model-key')
                    listener.settimeout(.05)
                    with self.assertRaises(TimeoutError):listener.accept()

    def test_app_close_reaps_running_sandbox_and_reopen_does_not_restart_it(self):
        with tempfile.TemporaryDirectory() as directory, FakeModelEndpoint('openai-responses','artificial-close',tool_name='python_analysis',tool_arguments={'code':'import time;time.sleep(20)'}) as model:
            root=Path(directory);secrets=MemorySecrets();database=root/'app.sqlite3'
            with TestClient(create_app(database,delay=0,credential_store=secrets)) as client:
                url,chat,identity,_,_,_=self.setup(client,root,model)
                state=wait_state(client,url+f'/conversations/{chat}',lambda s:s['tasks'][0]['toolExecutions'] and s['tasks'][0]['toolExecutions'][0].get('pythonPid'))
                pid=state['tasks'][0]['toolExecutions'][0]['pythonPid']
            with self.assertRaises(ProcessLookupError):os.kill(pid,0)
            with TestClient(create_app(database,delay=0,credential_store=secrets)) as client:
                reopened=client.post('/api/workbench/projects/open',json={'path':str(root/'project')}).json()
                self.assertEqual(reopened['tasks'][0]['status'],'paused')
                calls=len(model.requests);time.sleep(.1)
                self.assertEqual(len(model.requests),calls)

    def test_pause_and_node_kill_reap_analysis_then_selected_resume_only_redoes_unfinished(self):
        for operation in ('pause','kill'):
            with self.subTest(operation=operation),tempfile.TemporaryDirectory() as directory:
                root=Path(directory)
                slow="import time\nfrom pathlib import Path\nPath('started.txt').write_text('started')\ntime.sleep(3)\n"+SCRIPT
                with FakeModelEndpoint('openai-responses','artificial-pause',tool_name='python_analysis',tool_arguments={'code':slow}) as model:
                    database=root/'app.sqlite3';secrets=MemorySecrets()
                    with TestClient(create_app(database,delay=0,credential_store=secrets)) as client:
                        url,chat,identity,_,_,opened=self.setup(client,root,model)
                        state=wait_state(client,url+f'/conversations/{chat}',lambda s:s['tasks'][0]['toolExecutions'] and s['tasks'][0]['toolExecutions'][0].get('pythonPid'))
                        pid=state['tasks'][0]['toolExecutions'][0]['pythonPid']
                        if operation=='pause':self.assertEqual(client.post(url+f'/tasks/{identity}/pause').status_code,200)
                        else:os.kill(opened['runtimePid'],signal.SIGKILL)
                        deadline=time.monotonic()+5
                        while time.monotonic()<deadline:
                            try:os.kill(pid,0)
                            except ProcessLookupError:break
                            time.sleep(.02)
                        else:self.fail('sandbox child survived stopped Node')
                    with TestClient(create_app(database,delay=0,credential_store=secrets)) as client:
                        reopened=client.post('/api/workbench/projects/open',json={'path':str(root/'project')}).json()
                        self.assertEqual(reopened['tasks'][0]['status'],'paused')
                        count=reopened['modelRequests'];time.sleep(.1)
                        self.assertEqual(client.get(url).json()['modelRequests'],count)
                        client.put(url+f'/conversations/{chat}/model',json={'configurationId':None})
                        client.post(url+f'/tasks/{identity}/resume')
                        done=wait_task(client,url,chat,identity)
                        self.assertEqual(done['toolExecutions'][-1]['status'],'completed',done)
                        self.assertEqual(len(done['toolExecutions']),2)


if __name__=='__main__':unittest.main()
