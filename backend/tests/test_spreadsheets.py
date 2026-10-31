import io
from pathlib import Path
import tempfile
import time
import unittest
import zipfile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app_fixture import create_app
from silicolab.features.codon.spreadsheets import SpreadsheetError, inspect_table, synthesis_workbook

ROOT = Path(__file__).resolve().parent
TEMPLATE = Path(__file__).resolve().parents[1] / 'src/silicolab/features/codon/resources/synthesis.xlsx'


def workbook_bytes(rows, title='基因合成'):
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    content = io.BytesIO()
    book.save(content)
    book.close()
    return content.getvalue()


class SpreadsheetTests(unittest.TestCase):
    def test_synthetic_input_workbook_preserves_names_and_full_proteins(self):
        records = [[index, f'synthetic-{index}', 135, 'M' + 'K' * 134 + '*'] for index in range(1, 81)]
        content = workbook_bytes([['序号', '合成名', '长度', '序列'], *records])
        result = inspect_table(content, 'synthetic.xlsx')
        self.assertEqual(result['sheets'][0]['name'], '基因合成')
        rows = result['sheets'][0]['rows']
        self.assertEqual(rows[0], ['序号', '合成名', '长度', '序列'])
        self.assertEqual(len(rows) - 1, 80)
        self.assertEqual(rows[1][1], 'synthetic-1')
        self.assertEqual(len(rows[1][3]), 136)
        self.assertTrue(all(row[3].endswith('*') for row in rows[1:]))

    def test_multiple_sheets_blank_rows_extra_columns_and_headerless_data(self):
        book = Workbook()
        book.active.title = '说明'
        book.active.append(['Please select data sheet'])
        data = book.create_sheet('数据')
        for row in [['1', 'first', '', ' MK T* '], [None], ['2', 'second', 'note', 'MEF']]:
            data.append(row)
        content = io.BytesIO()
        book.save(content)
        book.close()
        result = inspect_table(content.getvalue(), 'test.xlsx')
        self.assertEqual([sheet['name'] for sheet in result['sheets']], ['说明', '数据'])
        self.assertEqual(result['sheets'][1]['rows'][0], ['1', 'first', '', ' MK T* '])
        self.assertEqual(result['sheets'][1]['rows'][1], [])
        self.assertEqual(result['sheets'][1]['rows'][2][-1], 'MEF')

    def test_formula_is_not_executed(self):
        result = inspect_table(workbook_bytes([['name', 'protein'], ['=1+1', 'MK']]), 'test.xlsx')
        self.assertEqual(result['sheets'][0]['rows'][1][0], '')

    def test_csv_tsv_quoting_and_chinese_encoding(self):
        result = inspect_table('备注,序列,合成名\r\n"带,逗号",MKT*,蛋白一\r\n'.encode('utf-8-sig'), 'test.csv')
        self.assertEqual(result['sheets'][0]['rows'][1], ['带,逗号', 'MKT*', '蛋白一'])
        result = inspect_table('合成名\t序列\n蛋白二\tMEF\n'.encode('gb18030'), 'test.tsv')
        self.assertEqual(result['sheets'][0]['rows'][1], ['蛋白二', 'MEF'])

    def test_rejects_lock_file_corrupt_legacy_empty_and_oversized_workbooks(self):
        for content, filename in [(b'bad', '~$test.xlsx'), (b'bad', 'test.xlsx'), (b'bad', 'test.xls'),
                                  (workbook_bytes([]), 'empty.xlsx'), (b'x' * (5 * 1024 * 1024 + 1), 'large.xlsx')]:
            with self.subTest(filename=filename), self.assertRaises(SpreadsheetError):
                inspect_table(content, filename)

    def test_rejects_decompressed_zip_bomb_and_too_many_columns(self):
        content = io.BytesIO()
        with zipfile.ZipFile(content, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('huge.xml', b'0' * (41 * 1024 * 1024))
        with self.assertRaisesRegex(SpreadsheetError, '解压'):
            inspect_table(content.getvalue(), 'bomb.xlsx')
        with self.assertRaisesRegex(SpreadsheetError, '列'):
            inspect_table(workbook_bytes([['name'] * 101]), 'wide.xlsx')

    def test_synthesis_export_matches_template_without_comments_or_extra_sheets(self):
        records = [
            {'name': 'first', 'sequence': 'MK*', 'original_sequence': 'MK', 'stop_added': True,
             'status': 'completed', 'dna': 'ATGAAGTAA', 'response': {'log_id': 'id-one'}},
            {'name': '=HYPERLINK("evil","click")', 'sequence': 'MK*', 'status': 'completed', 'dna': 'ATGAAGTAA'},
            {'name': 'failed-protein', 'sequence': 'MEF', 'status': 'failed', 'error': 'network error'},
        ]
        data = synthesis_workbook(records)
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ['基因合成'])
        synthesis = book['基因合成']
        headers = [cell.value for cell in synthesis[1]]
        self.assertEqual(headers, ['序号', '合成名', '核酸序列'])
        template = TEMPLATE
        if template.exists():
            original = load_workbook(template)
            reference = original.active
            self.assertEqual(headers, [cell.value for cell in next(reference.rows)])
            self.assertEqual(synthesis.title, reference.title)
            self.assertEqual(synthesis.freeze_panes, reference.freeze_panes)
            self.assertEqual(synthesis.auto_filter.ref, reference.auto_filter.ref)
            self.assertEqual(synthesis.sheet_view.showGridLines, reference.sheet_view.showGridLines)
            self.assertEqual(synthesis.sheet_format, reference.sheet_format)
            self.assertEqual(synthesis.page_setup, reference.page_setup)
            for column in ['A', 'B', 'C']:
                self.assertEqual(synthesis.column_dimensions[column].width, reference.column_dimensions[column].width)
                for row in [1, 2, 3]:
                    self.assertEqual(synthesis[f'{column}{row}']._style, reference[f'{column}{min(row, 2)}']._style)
            original.close()
        self.assertEqual(synthesis.max_row, 3)
        self.assertEqual(synthesis['B3'].data_type, 's')
        self.assertTrue(synthesis['B3'].value.startswith('='))
        self.assertEqual(synthesis['C2'].value, 'ATGAAGTAA')
        self.assertEqual(synthesis['A3'].value, 2)
        self.assertFalse(any(cell.comment is not None for sheet in book for row in sheet for cell in row))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertFalse(any('comment' in name.lower() for name in archive.namelist()))
        self.assertFalse(any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row))
        book.close()


class SpreadsheetApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.calls = []
        def optimize(payload):
            self.calls.append(payload)
            return {'success': '1', 'seq': ''.join({'M': 'ATG', 'K': 'AAG', '*': 'TAA'}[aa] for aa in payload['sequence'])}
        self.client = TestClient(create_app(Path(self.temp.name) / 'test.sqlite3', optimize, delay=0))
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.temp.cleanup()

    def wait(self, job_id):
        for _ in range(200):
            job = self.client.get('/api/jobs/' + job_id).json()
            if job['status'] not in {'queued', 'running'}:
                return job
            time.sleep(0.005)
        self.fail('job timed out')

    def test_import_returns_all_cells_without_optimizing(self):
        data = workbook_bytes([['notes', '序列', '合成名'], ['test', 'MK', 'protein']])
        response = self.client.post('/api/import/table?filename=test.xlsx', content=data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['sheets'][0]['rows'][1], ['test', 'MK', 'protein'])
        self.assertEqual(self.calls, [])
        self.assertEqual(self.client.post('/api/import/table?filename=bad.xlsx', content=b'bad').status_code, 400)
        self.assertEqual(self.client.post('/api/import/table?filename=big.xlsx', content=b'x' * (5 * 1024 * 1024 + 1)).status_code, 413)

    def test_add_stop_is_idempotent_and_retains_original_for_export(self):
        response = self.client.post('/api/jobs', json={'name': 'stop test', 'stop_policy': 'add',
            'records': [{'name': 'missing', 'sequence': 'MK'}, {'name': 'existing', 'sequence': 'MK*'}]})
        self.assertEqual(response.status_code, 201)
        job = self.wait(response.json()['id'])
        self.assertEqual(job['status'], 'completed')
        self.assertEqual([record['sequence'] for record in job['records']], ['MK*', 'MK*'])
        self.assertEqual([record['stop_added'] for record in job['records']], [True, False])
        self.assertEqual(job['records'][0]['original_sequence'], 'MK')
        self.assertEqual([call['sequence'] for call in self.calls], ['MK*', 'MK*'])
        exported = self.client.get(f"/api/jobs/{job['id']}/export?format=xlsx")
        self.assertEqual(exported.status_code, 200)
        self.assertIn('spreadsheetml', exported.headers['content-type'])
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book['基因合成']['C2'].value, 'ATGAAGTAA')
        self.assertEqual(book.sheetnames, ['基因合成'])
        self.assertEqual(book['基因合成']['C3'].value, 'ATGAAGTAA')
        self.assertFalse(any(cell.comment for sheet in book for row in sheet for cell in row))
        book.close()

    def test_preserve_stop_keeps_original_length(self):
        response = self.client.post('/api/jobs', json={'name': 'preserve', 'stop_policy': 'preserve', 'records': [{'name': 'a', 'sequence': 'MK'}]})
        job = self.wait(response.json()['id'])
        self.assertEqual(job['records'][0]['dna'], 'ATGAAG')
        self.assertFalse(job['records'][0]['stop_added'])
        self.assertFalse(job['records'][0]['checks']['stop'])

    def test_rejects_invalid_policy_internal_stop_and_oversize_addition_before_network(self):
        for sequence, policy in [('MK', 'unknown'), ('M*K', 'add'), ('M' * 10000, 'add')]:
            response = self.client.post('/api/jobs', json={'name': 'bad', 'stop_policy': policy, 'records': [{'name': 'a', 'sequence': sequence}]})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.calls, [])
        self.assertEqual(self.client.get('/api/jobs').json(), [])


if __name__ == '__main__':
    unittest.main()
