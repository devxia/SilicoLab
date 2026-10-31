"""Explicit ProteinFactory remote path and uncertain effect, injectable loopback only."""
import base64
import io
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import os
import signal
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import zipfile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from silicolab.server import create_app
from test_workbench_science import wait_task
from test_workbench_recovery import wait_state


class RemoteEndpoint:
    def __init__(self, mode):
        self.requests = []
        self.mode = mode
        self.release = threading.Event()
        endpoint = self
        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *_):
                pass
            def do_POST(self):
                data = json.loads(self.rfile.read(int(self.headers.get('Content-Length', 0))))
                endpoint.requests.append(data)
                protein = data['sequence']
                codons = {'M': 'ATG', 'K': 'AAG', 'E': 'GAG', 'F': 'TTT', '*': 'TAA'}
                response = {'success': '1', 'seq': ''.join(codons[aa] for aa in protein), 'log_id': 'artificial-remote'}
                mode = endpoint.mode
                if isinstance(mode, list):
                    mode = mode[min(len(endpoint.requests) - 1, len(mode) - 1)]
                if mode == 'accepted_hold':
                    endpoint.release.wait(6)
                if mode == 'invalid_dna':
                    response['seq'] = 'ATGTTTTAA'
                if mode == 'http_error':
                    self.send_error(503)
                    return
                if mode == 'accepted_then_drop':
                    self.close_connection = True
                    return
                raw = json.dumps(response).encode()
                self.send_response(200)
                self.send_header('Content-Type', 'application/json')
                self.send_header('Content-Length', str(len(raw)))
                self.end_headers()
                try:
                    self.wfile.write(raw)
                except (BrokenPipeError, ConnectionResetError):
                    pass
        self.server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        self.server.daemon_threads = True
        self.thread = __import__('threading').Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f'http://127.0.0.1:{self.server.server_address[1]}/toDnaAjaxCodon'

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()

    def close(self):
        self.release.set()
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)


class ExplicitRemoteTests(unittest.TestCase):
    def setup_task(self, client, root, source=b'>remote\nMK*\n', **overrides):
        project = client.post('/api/workbench/projects', json={'path': str(root)}).json()
        url = f'/api/workbench/projects/{project["id"]}'
        chat = client.post(url + '/conversations', json={'title': 'remote contract'}).json()['id']
        imported = client.post(url + '/files', json={'name': 'remote.fa', 'contentBase64': base64.b64encode(source).decode()}).json()
        fixed = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
        body = {'conversationId': chat, 'requestId': 'remote-contract', 'goal': '使用 ProteinFactory', 'workflow': 'protein_factory_remote', 'specialtyAcknowledged': True, 'inputVersionId': fixed['id'], 'authorizedTools': ['protein_factory_remote'], 'parameters': {'stopPolicy': 'preserve', 'avoid': ''}} | overrides
        response = client.post(url + '/tasks', json=body)
        self.assertEqual(response.status_code, 201, response.text)
        return url, chat, response.json()['id'], body

    def test_queued_retry_can_be_stopped_and_explicitly_resumed_without_losing_approval(self):
        for operation in ('pause', 'cancel', 'reopen'):
            with self.subTest(operation=operation), tempfile.TemporaryDirectory() as directory, RemoteEndpoint(['accepted_then_drop', 'accepted_hold', 'normal']) as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
                root = Path(directory)
                database = root / 'app.sqlite3'
                with TestClient(create_app(database, delay=0)) as client:
                    url, chat, task_id, body = self.setup_task(client, root)
                    target = url + f'/tasks/{task_id}'
                    wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] == 'result_unknown')
                    blocker = client.post(url + '/tasks', json=body | {'requestId': 'blocker'}).json()
                    wait_state(client, url + f'/conversations/{chat}', lambda s: len(remote.requests) == 2)
                    question = client.post(target + '/reconcile').json()['decisions'][-1]
                    approved = client.post(target + f'/decisions/{question["id"]}', json={'answer': 'retry_remote'}).json()
                    self.assertEqual(approved['status'], 'admission_pending')
                    self.assertEqual(client.post(target + '/reconcile').json()['status'], 'admission_pending')
                    self.assertEqual(client.post(target + '/continue').json()['status'], 'queued')
                    if operation != 'reopen':
                        stopped = client.post(target + '/' + operation).json()
                        self.assertEqual(stopped['status'], 'paused' if operation == 'pause' else 'cancelled')
                        remote.release.set()
                        wait_task(client, url, chat, blocker['id'])
                        self.assertEqual(len(remote.requests), 2)
                        if operation == 'pause':
                            self.assertEqual(client.post(target + '/reconcile').json()['status'], 'paused')
                            self.assertEqual(client.post(target + '/resume').status_code, 200)
                            done = wait_task(client, url, chat, task_id)
                            self.assertEqual(len(done['remoteAttempts']), 2)
                            self.assertEqual(len(remote.requests), 3)
                        else:
                            self.assertEqual(client.post(target + '/resume').json()['status'], 'cancelled')
                            self.assertEqual(len(remote.requests), 2)
                if operation == 'reopen':
                    remote.release.set()
                    with TestClient(create_app(database, delay=0)) as client:
                        reopened = client.post('/api/workbench/projects/open', json={'path': str(root)}).json()
                        self.assertEqual(reopened['tasks'][0]['status'], 'paused')
                        self.assertEqual(len(remote.requests), 2)
                        self.assertEqual(client.post(target + '/resume').status_code, 200)
                        done = wait_task(client, url, chat, task_id)
                        self.assertEqual(len(done['remoteAttempts']), 2)
                        self.assertEqual(len(remote.requests), 3)

    def test_tampered_partial_success_prevents_retry_before_new_http(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint(['normal', 'accepted_then_drop']) as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                url, chat, task_id, _ = self.setup_task(client, root, b'>first\nMK*\n>second\nME*\n')
                unknown = wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] == 'result_unknown')['tasks'][0]
                artifact = next(a for a in unknown['toolExecutions'][0]['artifacts'] if a['name'] == 'synthesis.fasta')
                (root / artifact['path']).write_text('tampered artificial result')
                target = url + f'/tasks/{task_id}'
                question = client.post(target + '/reconcile').json()['decisions'][-1]
                denied = client.post(target + f'/decisions/{question["id"]}', json={'answer': 'retry_remote'})
                self.assertEqual(denied.status_code, 409, denied.text)
                kept = client.post(target + f'/decisions/{question["id"]}', json={'answer': 'keep_unknown'})
                self.assertEqual(kept.status_code, 200, kept.text)
                self.assertEqual(kept.json()['status'], 'result_unknown')
                self.assertEqual(kept.json()['decisions'][-1]['answer'], 'keep_unknown')
                self.assertEqual(len(remote.requests), 2)

    def test_presend_validation_failure_is_failed_not_unknown_and_never_requests(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint('normal') as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                url, chat, task_id, _ = self.setup_task(client, root, parameters={'stopPolicy': 'preserve', 'avoid': 'XYZ'})
                done = wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] in ('failed', 'result_unknown'))['tasks'][0]
                self.assertEqual(done['status'], 'failed')
                self.assertEqual(done['toolExecutions'][0]['status'], 'failed')
                self.assertEqual(done['remoteAttempts'][0]['status'], 'failed')
                self.assertEqual(remote.requests, [])
                self.assertEqual(client.post(url + f'/tasks/{task_id}/reconcile').json()['status'], 'failed')
                self.assertFalse(done['decisions'])

    def test_explicit_retry_preserves_successful_records_and_is_consumed_once(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint(['normal', 'accepted_then_drop']) as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                url, chat, task_id, body = self.setup_task(client, root, b'>first\nMK*\n>second\nME*\n')
                unknown = wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] == 'result_unknown')['tasks'][0]
                self.assertEqual(len(remote.requests), 2)
                first = unknown['toolExecutions'][0]
                self.assertEqual([r['status'] for r in first['result']['records']], ['completed', 'unknown'])
                target = url + f'/tasks/{task_id}'
                checked = client.post(target + '/reconcile').json()
                question = checked['decisions'][-1]
                self.assertIn('可能重复', question['question'])
                remote.mode = 'normal'
                answer_url = target + f'/decisions/{question["id"]}'
                self.assertEqual(client.post(answer_url, json={'answer': 'retry_remote'}).json()['status'], 'admission_pending')
                self.assertEqual(client.post(answer_url, json={'answer': 'retry_remote'}).json()['status'], 'admission_pending')
                self.assertEqual(len(remote.requests), 2)
                self.assertEqual(client.post(target + '/continue').status_code, 200)
                self.assertEqual(client.post(target + '/continue').status_code, 200)
                done = wait_task(client, url, chat, task_id)
                self.assertEqual(len(remote.requests), 3)
                self.assertEqual(remote.requests[-1]['sequence'], 'ME*')
                self.assertEqual(done['toolExecutions'][0], first)
                self.assertTrue(done['toolExecutions'][1]['result']['records'][0]['reused'])
                self.assertNotEqual(done['remoteAttempts'][0]['id'], done['remoteAttempts'][1]['id'])
                self.assertEqual(client.post(target + '/reconcile').json(), done)
                self.assertEqual(client.post(url + '/tasks', json=body).json(), done)
                self.assertEqual(len(remote.requests), 3)

    def test_invalid_success_and_http_errors_do_not_become_scientific_success(self):
        for mode in ('invalid_dna', 'http_error'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory, RemoteEndpoint(mode) as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
                root = Path(directory)
                with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                    url, chat, task_id, _ = self.setup_task(client, root)
                    done = wait_task(client, url, chat, task_id)
                    execution = done['toolExecutions'][0]
                    self.assertEqual(execution['status'], 'failed')
                    self.assertEqual(execution['result']['status'], 'failed')
                    self.assertNotIn('dna', execution['result']['records'][0])
                    self.assertFalse(any(a['name'] == 'synthesis.fasta' for a in execution['artifacts']))
                    self.assertEqual(len(remote.requests), 1)

    def test_true_timeout_and_pausing_inflight_remote_never_authorize_replay(self):
        for operation in ('timeout', 'pause', 'cancel'):
            with self.subTest(operation=operation), tempfile.TemporaryDirectory() as directory, RemoteEndpoint('accepted_hold') as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url, 'SILICOLAB_SCIENCE_REMOTE_TIMEOUT': '0.2' if operation == 'timeout' else '10'}):
                root = Path(directory)
                database = root / 'app.sqlite3'
                with TestClient(create_app(database, delay=0)) as client:
                    url, chat, task_id, _ = self.setup_task(client, root)
                    wait_state(client, url + f'/conversations/{chat}', lambda s: len(remote.requests) == 1)
                    target = url + f'/tasks/{task_id}'
                    if operation != 'timeout':
                        stopped = client.post(target + '/' + operation)
                        self.assertEqual(stopped.status_code, 200, stopped.text)
                        self.assertEqual(stopped.json()['status'], 'result_unknown')
                    unknown = wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] == 'result_unknown')['tasks'][0]
                    self.assertEqual(unknown['toolExecutions'][0]['status'], 'unknown')
                    self.assertEqual(client.post(target + '/continue').json()['status'], 'result_unknown')
                    self.assertEqual(client.post(target + '/resume').json()['status'], 'result_unknown')
                    self.assertEqual(client.post(target + '/cancel').json()['status'], 'result_unknown')
                    self.assertEqual(len(remote.requests), 1)
                with TestClient(create_app(database, delay=0)) as client:
                    reopened = client.post('/api/workbench/projects/open', json={'path': str(root)}).json()
                    self.assertEqual(reopened['tasks'][0]['status'], 'result_unknown')
                    self.assertEqual(len(remote.requests), 1)

    def test_authorization_is_task_scoped_and_regular_sequence_requests_never_upload(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint('normal') as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                url, chat, task_id, _ = self.setup_task(client, root, authorizedTools=[])
                task = client.get(url).json()['tasks'][0]
                self.assertEqual(task['status'], 'needs_decision')
                self.assertEqual(remote.requests, [])
                question = task['decisions'][0]
                answer = client.post(url + f'/tasks/{task_id}/decisions/{question["id"]}', json={'answer': 'grant'})
                self.assertEqual(answer.status_code, 200, answer.text)
                self.assertEqual(answer.json()['authorizedTools'], ['protein_factory_remote'])
                self.assertEqual(remote.requests, [])
                client.post(url + f'/tasks/{task_id}/continue')
                wait_task(client, url, chat, task_id)
                ordinary = client.post(url + '/tasks', json={'conversationId': chat, 'requestId': 'ordinary', 'goal': '普通蛋白转 DNA，不使用专项', 'authorizedTools': ['synthetic_count']}).json()
                wait_task(client, url, chat, ordinary['id'])
                self.assertEqual(len(remote.requests), 1)
                denied = client.post(url + '/tasks', json={'conversationId': chat, 'requestId': 'wrong-scope', 'goal': '普通序列', 'authorizedTools': ['protein_factory_remote']})
                self.assertEqual(denied.status_code, 409)
                self.assertEqual(len(remote.requests), 1)

    def test_host_kill_reaps_remote_worker_and_reopen_does_not_resubmit(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint('accepted_hold') as remote, patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
            root = Path(directory)
            database = root / 'app.sqlite3'
            with TestClient(create_app(database, delay=0)) as client:
                url, chat, task_id, _ = self.setup_task(client, root)
                state = wait_state(client, url + f'/conversations/{chat}', lambda s: len(remote.requests) == 1 and s['tasks'][0]['toolExecutions'][0].get('pythonPid'))
                pid = state['tasks'][0]['toolExecutions'][0]['pythonPid']
                os.kill(client.get(url).json()['runtimePid'], signal.SIGKILL)
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    try:
                        os.kill(pid, 0)
                    except ProcessLookupError:
                        break
                    time.sleep(0.02)
                else:
                    self.fail('Remote worker survived its Node owner')
                reopened = client.post('/api/workbench/projects/open', json={'path': str(root)}).json()
                self.assertEqual(reopened['tasks'][0]['status'], 'result_unknown')
                self.assertEqual(reopened['tasks'][0]['toolExecutions'][0]['status'], 'unknown')
                checked = client.post(url + f'/tasks/{task_id}/reconcile').json()
                self.assertEqual(checked['status'], 'needs_decision')
                self.assertFalse(checked['remoteAttempts'][0]['queryAvailable'])
                self.assertTrue(checked['remoteAttempts'][0]['reconciledAt'])
                self.assertEqual(len(remote.requests), 1)

    def test_explicit_remote_validates_and_exports_real_response_without_fallback(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint('normal') as remote:
            root = Path(directory)
            with patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
                with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                    project = client.post('/api/workbench/projects', json={'path': str(root)}).json()
                    url = f'/api/workbench/projects/{project["id"]}'
                    chat = client.post(url + '/conversations', json={'title': 'remote explicit'}).json()['id']
                    source = b'>remote\nMK*\n'
                    imported = client.post(url + '/files', json={'name': 'remote.fa', 'contentBase64': base64.b64encode(source).decode()}).json()
                    fixed = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                    response = client.post(url + '/tasks', json={'conversationId': chat, 'requestId': 'remote-1', 'goal': '使用 ProteinFactory 远端', 'authorizedTools': ['protein_factory_remote'], 'workflow': 'protein_factory_remote', 'specialtyAcknowledged': True, 'inputVersionId': fixed['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'GAATTC'}})
                    self.assertEqual(response.status_code, 201, response.text)
                    task = wait_task(client, url, chat, response.json()['id'])
                    execution = task['toolExecutions'][0]
                    self.assertEqual(execution['result']['records'][0]['dna'], 'ATGAAGTAA')
                    self.assertEqual(execution['result']['engine_info']['id'], 'remote')
                    self.assertEqual(execution['result']['network_attempts'], 1)
                    self.assertEqual(len(remote.requests), 1)
                    self.assertEqual(remote.requests[0]['speciesName'], 'D2P(default)')
                    self.assertEqual(remote.requests[0]['sequence'], 'MK*')
                    self.assertEqual(client.get('/api/jobs').json(), [])
                    artifact_url = url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}'
                    fasta = client.get(artifact_url + '/synthesis.fasta')
                    self.assertEqual(fasta.text, '>remote\nATGAAGTAA\n')
                    excel = client.get(artifact_url + '/synthesis.xlsx')
                    workbook = load_workbook(io.BytesIO(excel.content))
                    self.assertEqual(workbook.sheetnames, ['基因合成'])
                    self.assertEqual(list(workbook.active.values)[1], (1, 'remote', 'ATGAAGTAA'))
                    self.assertTrue(all(cell.comment is None for row in workbook.active for cell in row))
                    workbook.close()
                    with zipfile.ZipFile(io.BytesIO(client.get(artifact_url + '/records.zip').content)) as archive:
                        saved = json.loads(archive.read('record.json'))
                        self.assertEqual(saved['records'][0]['raw_response']['seq'], 'ATGAAGTAA')
                        self.assertEqual(saved['records'][0]['source_file_id'], imported['id'])

    def test_timeout_after_accept_is_unknown_never_fallback_or_automatic_retry(self):
        with tempfile.TemporaryDirectory() as directory, RemoteEndpoint('accepted_then_drop') as remote:
            root = Path(directory)
            with patch.dict('os.environ', {'SILICOLAB_SCIENCE_REMOTE': remote.url}):
                with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                    project = client.post('/api/workbench/projects', json={'path': str(root)}).json()
                    url = f'/api/workbench/projects/{project["id"]}'
                    chat = client.post(url + '/conversations', json={'title': 'uncertain science'}).json()['id']
                    imported = client.post(url + '/files', json={'name': 'uncertain.fa', 'contentBase64': base64.b64encode(b'>uncertain\nMK*\n').decode()}).json()
                    fixed = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                    body = {'conversationId': chat, 'requestId': 'uncertain-1', 'goal': '使用 ProteinFactory 远端', 'authorizedTools': ['protein_factory_remote'], 'workflow': 'protein_factory_remote', 'specialtyAcknowledged': True, 'inputVersionId': fixed['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': ''}}
                    response = client.post(url + '/tasks', json=body)
                    self.assertEqual(response.status_code, 201, response.text)
                    task = wait_state(client, url + f'/conversations/{chat}', lambda s: s['tasks'][0]['status'] == 'result_unknown')['tasks'][0]
                    self.assertEqual(task['status'], 'result_unknown')
                    self.assertEqual(task['toolExecutions'][0]['status'], 'unknown')
                    checked = client.post(url + f'/tasks/{task["id"]}/reconcile').json()
                    self.assertEqual(checked['status'], 'needs_decision')
                    question = next(item for item in checked['decisions'] if item['kind'] == 'science_retry')
                    self.assertTrue(question['uncertainty']['transportTimeout'])
                    self.assertFalse(question['uncertainty']['queryAvailable'])
                    self.assertEqual(len(remote.requests), 1)
                    self.assertEqual(client.get('/api/jobs').json(), [])
                    answer_url = url + f'/tasks/{task["id"]}/decisions/{question["id"]}'
                    preserved = client.post(answer_url, json={'answer': 'keep_unknown'}).json()
                    self.assertEqual(preserved['status'], 'result_unknown')
                    self.assertEqual(client.post(answer_url, json={'answer': 'keep_unknown'}).json()['status'], 'result_unknown')
                    self.assertEqual(len(remote.requests), 1)
                    self.assertEqual(client.post(url + f'/tasks/{task["id"]}/continue').json()['status'], 'result_unknown')
                    self.assertFalse(any(artifact['name'] == 'synthesis.fasta' for artifact in task['toolExecutions'][0].get('artifacts', [])))


if __name__ == '__main__':
    unittest.main()
