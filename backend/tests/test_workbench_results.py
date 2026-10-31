"""Partial scientific results and complete exports match actual task evidence."""
import base64
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app_fixture import create_app
from test_workbench_science import wait_task


class ScientificResultTests(unittest.TestCase):
    def test_partial_results_keep_order_individual_fasta_and_complete_record_bundle(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                project = client.post('/api/workbench/projects', json={'path': str(root)}).json()
                url = f'/api/workbench/projects/{project["id"]}'
                chat = client.post(url + '/conversations', json={'title': 'partial records'}).json()['id']
                source = b'>good-1\nK*\n>bad\nM*\n>good-2\nKK*\n'
                imported = client.post(url + '/files', json={'name': 'mixed.fasta', 'contentBase64': base64.b64encode(source).decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                task = client.post(url + '/tasks', json={'conversationId': chat, 'requestId': 'partial', 'goal': 'D2P 专项本地实验并核对部分失败', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'ATG'}}).json()
                completed = wait_task(client, url, chat, task['id'])
                execution = completed['toolExecutions'][0]
                report = execution['result']
                self.assertEqual(report['status'], 'partial')
                self.assertEqual([record['status'] for record in report['records']], ['completed', 'failed', 'completed'])
                self.assertEqual([record['name'] for record in report['records']], ['good-1', 'bad', 'good-2'])
                self.assertTrue(report['records'][1]['error'])
                self.assertTrue(report['records'][0]['raw_response'])
                self.assertEqual(report['attempt']['id'], execution['id'])
                artifact_url = url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}'
                individual = client.get(artifact_url + '/sequence-001.fasta')
                self.assertEqual(individual.status_code, 200, individual.text)
                self.assertEqual(individual.text, '>good-1\nAAGTAA\n')
                self.assertEqual(client.get(artifact_url + '/sequence-002.fasta').status_code, 409)
                bulk = client.get(artifact_url + '/synthesis.fasta')
                self.assertEqual(bulk.text, '>good-1\nAAGTAA\n>good-2\nAAGAAGTAA\n')
                excel = client.get(artifact_url + '/synthesis.xlsx')
                workbook = load_workbook(io.BytesIO(excel.content))
                self.assertEqual(workbook.sheetnames, ['基因合成'])
                self.assertEqual(list(workbook.active.values), [('序号', '合成名', '核酸序列'), (1, 'good-1', 'AAGTAA'), (2, 'good-2', 'AAGAAGTAA')])
                self.assertTrue(all(cell.comment is None for row in workbook.active for cell in row))
                workbook.close()
                bundle = client.get(artifact_url + '/records.zip')
                self.assertEqual(bundle.status_code, 200, bundle.text[:100])
                with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
                    saved = json.loads(archive.read('record.json'))
                    self.assertEqual(saved, report)
                    self.assertEqual(json.loads(archive.read('task.json'))['id'], task['id'])
                    self.assertEqual(json.loads(archive.read('input-version.json'))['id'], version['id'])
                    self.assertEqual(archive.read('inputs/001.fasta'), source)
                    self.assertEqual(archive.read('synthesis.xlsx'), excel.content)
                    self.assertEqual(archive.read('synthesis.fasta'), bulk.content)
                    self.assertEqual(saved['records'][0]['raw_response']['seq'], 'AAGTAA')
                current = client.get(url + f'/conversations/{chat}').json()
                self.assertEqual(len(current['tasks'][0]['toolExecutions']), 1)
                self.assertEqual(current['tasks'][0]['toolExecutions'][0], execution)
                # Tamper a saved artifact: download must fail rather than return unchecked bytes.
                stored = next(item for item in execution['artifacts'] if item['name'] == 'records.zip')
                (root / stored['path']).write_bytes(b'tampered')
                self.assertEqual(client.get(artifact_url + '/records.zip').status_code, 409)


if __name__ == '__main__':
    unittest.main()
