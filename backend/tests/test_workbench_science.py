"""Explicit local D2P task: fixed input -> real Node/Python -> verified DNA/download."""
import base64
import io
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from silicolab.features.codon.client import validate_result
from silicolab.server import create_app


def wait_task(client, url, chat, identity):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        state = client.get(url + f'/conversations/{chat}').json()
        task = next(item for item in state['tasks'] if item['id'] == identity)
        if task['status'] in ['answered', 'failed', 'cancelled']:
            return task
        time.sleep(0.03)
    raise AssertionError(state)


class LocalScienceTests(unittest.TestCase):
    def test_explicit_local_task_reuses_algorithm_validates_and_downloads_real_artifacts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                identity = client.post('/api/workbench/projects', json={'path': str(root)}).json()['id']
                url = f'/api/workbench/projects/{identity}'
                chat = client.post(url + '/conversations', json={'title': 'D2P local'}).json()['id']
                source = b'>first\nMK*\n>second\nMEFEFK*\n'
                imported = client.post(url + '/files', json={'name': 'proteins.fasta', 'contentBase64': base64.b64encode(source).decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                body = {'conversationId': chat, 'requestId': 'local-explicit', 'goal': '使用 ProteinFactory / D2P 专项本地实验，保留终止符并校验导出', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'GAATTC'}}
                response = client.post(url + '/tasks', json=body)
                self.assertEqual(response.status_code, 201, response.text)
                task = wait_task(client, url, chat, response.json()['id'])
                self.assertEqual(client.post(url + '/tasks', json={**body, 'specialtyAcknowledged': False}).status_code, 409)
                without_ack = dict(body)
                del without_ack['specialtyAcknowledged']
                self.assertEqual(client.post(url + '/tasks', json=without_ack).status_code, 409)
                execution = task['toolExecutions'][0]
                self.assertEqual(execution['status'], 'completed', execution)
                result = execution['result']
                self.assertEqual(result['engine_info']['id'], 'local')
                self.assertFalse(result['engine_info']['network_used'])
                self.assertEqual(result['network_attempts'], 0)
                self.assertEqual(result['network_guard_self_test'], 'blocked')
                self.assertEqual(result['inputVersionId'], version['id'])
                self.assertEqual(len(result['records']), 2)
                for record in result['records']:
                    self.assertEqual(validate_result({'success': '1', 'seq': record['dna']}, record['protein'], 'GAATTC'), record['dna'])
                    self.assertTrue(record['checks']['translation'])
                    self.assertTrue(record['protein'].endswith('*'))
                artifacts = execution['artifacts']
                fasta = client.get(url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}/synthesis.fasta')
                self.assertEqual(fasta.status_code, 200, fasta.text)
                self.assertIn('>first', fasta.text)
                self.assertIn(result['records'][0]['dna'], fasta.text)
                excel = client.get(url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}/synthesis.xlsx')
                workbook = load_workbook(io.BytesIO(excel.content))
                self.assertEqual(workbook.sheetnames, ['基因合成'])
                self.assertEqual(list(workbook.active.values)[0], ('序号', '合成名', '核酸序列'))
                self.assertTrue(all(cell.comment is None for row in workbook.active for cell in row))
                self.assertTrue(all((root / item['path']).is_file() for item in artifacts))
                self.assertEqual(client.get('/api/jobs').json(), [])
                # Ordinary goals must never infer this specialty from keywords alone.
                ordinary = client.post(url + '/tasks', json={'conversationId': chat, 'requestId': 'ordinary', 'goal': '优化密码子', 'authorizedTools': ['synthetic_count']})
                plain = wait_task(client, url, chat, ordinary.json()['id'])
                self.assertEqual([item['tool'] for item in plain['toolExecutions']], ['synthetic_count'])

    def test_infeasible_and_invalid_local_requests_never_fallback_or_claim_success(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                identity = client.post('/api/workbench/projects', json={'path': str(root)}).json()['id']
                url = f'/api/workbench/projects/{identity}'
                chat = client.post(url + '/conversations', json={'title': 'infeasible'}).json()['id']
                imported = client.post(url + '/files', json={'name': 'one.fa', 'contentBase64': base64.b64encode(b'>only\nM*\n').decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                body = {'conversationId': chat, 'requestId': 'infeasible', 'goal': '明确使用 D2P 专项本地实验', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'ATG'}}
                self.assertEqual(client.post(url + '/tasks', json={**body, 'specialtyAcknowledged': False}).status_code, 409)
                task = client.post(url + '/tasks', json=body).json()
                finished = wait_task(client, url, chat, task['id'])
                execution = finished['toolExecutions'][0]
                self.assertEqual(execution['status'], 'failed')
                self.assertEqual(execution['result']['status'], 'failed')
                self.assertEqual(execution['result']['network_attempts'], 0)
                self.assertEqual(execution['result']['records'][0]['status'], 'failed')
                self.assertEqual([artifact['name'] for artifact in execution['artifacts']], ['records.zip', 'record.json'])
                self.assertEqual(client.get('/api/jobs').json(), [])
                self.assertNotIn('dna', execution['result']['records'][0])

    def test_cancel_reports_actual_python_termination(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                identity = client.post('/api/workbench/projects', json={'path': str(root)}).json()['id']
                url = f'/api/workbench/projects/{identity}'
                chat = client.post(url + '/conversations', json={'title': 'cancel local'}).json()['id']
                source = ('>long\nM' + 'K' * 8000 + '*\n').encode()
                imported = client.post(url + '/files', json={'name': 'long.fa', 'contentBase64': base64.b64encode(source).decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                body = {'conversationId': chat, 'requestId': 'cancel', 'goal': 'D2P 专项本地实验', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'AAG'}}
                task = client.post(url + '/tasks', json=body).json()
                deadline = time.monotonic() + 8
                pid = None
                while time.monotonic() < deadline:
                    state = client.get(url + f'/conversations/{chat}').json()
                    attempts = state['tasks'][0]['toolExecutions']
                    if attempts and attempts[0].get('pythonPid'):
                        pid = attempts[0]['pythonPid']
                        break
                    time.sleep(0.005)
                self.assertIsNotNone(pid, state)
                cancelled = client.post(url + f'/tasks/{task["id"]}/cancel')
                self.assertEqual(cancelled.status_code, 200, cancelled.text)
                self.assertEqual(cancelled.json()['status'], 'cancelled')
                with self.assertRaises(ProcessLookupError):
                    os.kill(pid, 0)
                self.assertTrue(all(attempt['status'] != 'running' for attempt in cancelled.json()['toolExecutions']))
                self.assertEqual(client.get('/api/jobs').json(), [])

    def test_stop_presence_is_separate_from_policy_check_and_preparation_failure_preserves_success(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with TestClient(create_app(root / 'app.sqlite3', delay=0)) as client:
                identity = client.post('/api/workbench/projects', json={'path': str(root)}).json()['id']
                url = f'/api/workbench/projects/{identity}'
                chat = client.post(url + '/conversations', json={'title': 'stop policies'}).json()['id']
                imported = client.post(url + '/files', json={'name': 'no-stop.fa', 'contentBase64': base64.b64encode(b'>short\nMK\n').decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                body = {'conversationId': chat, 'requestId': 'preserve-no-stop', 'goal': '明确 D2P 专项本地实验', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': ''}}
                task = client.post(url + '/tasks', json=body).json()
                preserved = wait_task(client, url, chat, task['id'])['toolExecutions'][0]['result']['records'][0]
                self.assertEqual(preserved['status'], 'completed')
                self.assertEqual(preserved['protein'], 'MK')
                self.assertFalse(preserved['terminal_stop_present'])
                self.assertTrue(preserved['checks']['stop_policy'])
                body['requestId'] = 'add-stop'
                body['parameters']['stopPolicy'] = 'add'
                task = client.post(url + '/tasks', json=body).json()
                added = wait_task(client, url, chat, task['id'])['toolExecutions'][0]['result']['records'][0]
                self.assertEqual(added['protein'], 'MK*')
                self.assertTrue(added['terminal_stop_present'])
                self.assertTrue(added['checks']['stop_policy'])
                mixed = b'>good\nMK*\n>too-long-after-add\n' + b'K' * 10000 + b'\n'
                imported = client.post(url + '/files', json={'name': 'mixed.fa', 'contentBase64': base64.b64encode(mixed).decode()}).json()
                version = client.post(url + '/input-versions', json={'conversationId': chat, 'fileIds': [imported['id']]}).json()
                body.update(requestId='mixed-preparation', inputVersionId=version['id'])
                task = client.post(url + '/tasks', json=body).json()
                finished = wait_task(client, url, chat, task['id'])
                execution = finished['toolExecutions'][0]
                self.assertEqual(execution['result']['status'], 'partial', execution)
                records = execution['result']['records']
                self.assertEqual(records[0]['status'], 'completed')
                self.assertEqual(records[0]['dna'], 'ATGAAGTAA')
                self.assertEqual(records[1]['status'], 'failed')
                self.assertIn('10,000', records[1]['error'])
                fasta = client.get(url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}/synthesis.fasta')
                self.assertEqual(fasta.status_code, 200, fasta.text)
                self.assertEqual(fasta.text, '>good\nATGAAGTAA\n')
                excel = client.get(url + f'/tasks/{task["id"]}/artifacts/{execution["id"]}/synthesis.xlsx')
                workbook = load_workbook(io.BytesIO(excel.content))
                self.assertEqual(list(workbook.active.values)[1:], [(1, 'good', 'ATGAAGTAA')])
                workbook.close()
                self.assertEqual(execution['result']['network_attempts'], 0)


if __name__ == '__main__':
    unittest.main()
