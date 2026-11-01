"""Audit a built product inventory, secret/private-path exclusions and shipped resources."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess
import zipfile


def audit(app,manifest,source_root):
    app=Path(app).resolve();inventory=json.loads(Path(manifest).read_text());source_root=Path(source_root).resolve()
    subprocess.run(['/usr/bin/codesign','--verify','--deep','--strict',str(app)],check=True,capture_output=True)
    for name,expected in inventory['files'].items():
        path=app/name
        if not path.is_file() or path.is_symlink() or hashlib.sha256(path.read_bytes()).hexdigest()!=expected:raise ValueError('Bundle inventory mismatch: '+name)
    for name,target in inventory['links'].items():
        path=app/name
        if not path.is_symlink() or str(path.readlink())!=target or not path.resolve().is_relative_to(app):raise ValueError('Bundle link mismatch: '+name)
    actual_files={p.relative_to(app).as_posix() for p in app.rglob('*') if p.is_file() and not p.is_symlink()}
    actual_links={p.relative_to(app).as_posix() for p in app.rglob('*') if p.is_symlink()}
    if actual_files!=set(inventory['files']) or actual_links!=set(inventory['links']):raise ValueError('Application contains files or links absent from the inventory')
    excluded={'research','private-inputs','outputs','.git','.venv','.desktop-venv','.silicolab'}
    disallowed_files={'auth.json','model-configs.json','projects.json','notification-state.json','mcp-trust.json','background.json','gui-session.json','workbench-selection.json'}
    local_markers=[str(source_root).encode(),str(Path.home()).encode()]
    credential=re.compile(rb'(?i)(?<![A-Za-z0-9_-])(?:sk-[A-Za-z0-9_-]{30,}|gh[pousr]_[A-Za-z0-9]{30,})')
    checked=0
    for path in app.rglob('*'):
        if path.is_symlink() or not path.is_file():continue
        relative=path.relative_to(app)
        if excluded.intersection(relative.parts) or path.name in disallowed_files or path.suffix in ('.sqlite','.sqlite3','.log','.pid'):raise ValueError('Runtime/private content in bundle: '+str(relative))
        data=path.read_bytes();checked+=1
        if any(marker in data for marker in local_markers):raise ValueError('Local source/home path in bundle bytes: '+str(relative))
        if credential.search(data):raise ValueError('Credential-like bytes require manual review: '+str(relative))
    # Inspect decompressed frozen code too: raw executable scans cannot see PYZ constants.
    from PyInstaller.archive.readers import CArchiveReader
    frozen=CArchiveReader(str(app/'Contents/MacOS/SilicoLab'));modules=0
    for name,entry in frozen.toc.items():
        if entry[-1]=='z':
            embedded=frozen.open_embedded_archive(name)
            for module in embedded.toc:
                data=embedded.extract(module,raw=True)
                if data:
                    modules+=1
                    if any(marker in data for marker in local_markers):raise ValueError('Local build path in frozen module: '+module)
                    if module.startswith('silicolab') and credential.search(data):raise ValueError('Credential-like bytes in app code')
        elif entry[-1] in ('s','m','M'):
            data=frozen.extract(name)
            if any(marker in data for marker in local_markers):raise ValueError('Local build path in executable entry: '+name)
    for archive in app.rglob('*.zip'):
        if archive.is_symlink():continue
        with zipfile.ZipFile(archive) as zipped:
            if sum(i.file_size for i in zipped.infolist())>256*1024*1024:raise ValueError('Embedded archive exceeds audit bound')
            for info in zipped.infolist():
                if info.is_dir():continue
                data=zipped.read(info)
                if any(marker in data for marker in local_markers):raise ValueError('Local build path in embedded archive entry: '+info.filename)
    resources=app/'Contents/Resources'
    if (resources/'analysis/PythonBoundaryProbe.zip').exists():raise ValueError('Redundant analysis runtime ZIP must remain outside the product')
    required=['runtime/bin/node','runtime/NODE-LICENSE','runtime/package-lock.json','runtime/node_modules/@earendil-works/pi-durable/package.json','runtime/node_modules/@earendil-works/pi-ai/package.json','analysis/PythonBoundaryProbe.app','analysis/PythonBoundaryProbe-manifest.json','notifications/SilicoLabNotifications.app','notifications/SilicoLabNotifications-manifest.json','frontend/dist/fonts/DM-Sans.ttf','frontend/dist/fonts/OFL.txt','certifi/cacert.pem']
    for name in required:
        if not (resources/name).exists():raise ValueError('Required shipped resource missing: '+name)
    templates=list((resources/'silicolab/features/codon/resources').glob('*.xlsx'))
    if len(templates)!=1:raise ValueError('Blank synthesis template missing')
    from openpyxl import load_workbook
    book=load_workbook(templates[0]);rows=[row for row in book.active.values if any(value is not None for value in row)]
    if book.sheetnames!=['基因合成'] or len(rows)!=1 or rows[0]!=('序号','合成名','核酸序列') or any(c.comment for row in book.active for c in row):raise ValueError('Template contains nonblank/private data')
    book.close()
    return {'status':'passed','regularFilesChecked':checked,'frozenModulesChecked':modules,'internalLinksChecked':len(inventory['links']),'requiredResources':len(required),'privateStateAbsent':True,'localSourceHomePathsAbsent':True,'credentialPatternScan':'clear','blankTemplate':True,'signature':'ad-hoc; not Developer ID or notarized'}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--app',type=Path,required=True);parser.add_argument('--manifest',type=Path,required=True);parser.add_argument('--source-root',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    result=audit(args.app,args.manifest,args.source_root);args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result))
