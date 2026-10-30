"""Verify local UI execution with ALL server-side external network calls forbidden."""
import io
import json
from pathlib import Path
import zipfile
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

URL = 'http://127.0.0.1:8766'
OUT = Path('outputs/ui'); OUT.mkdir(parents=True, exist_ok=True)

with sync_playwright() as p:
    browser = p.chromium.launch(channel='chrome', headless=True)
    context = browser.new_context(viewport={'width': 1440, 'height': 1060})
    external_requests = []
    page = context.new_page()
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.on('request', lambda request: external_requests.append(request.url) if not request.url.startswith(URL) else None)
    page.goto(URL + '/#manual', wait_until='networkidle')
    expect(page.get_by_role('radio', name='ProteinFactory 远端', exact=True)).to_be_checked()
    page.get_by_label('目标宿主', exact=True).select_option('Escherichia_coli')
    page.get_by_role('radio', name='本地实验', exact=True).check()
    expect(page.get_by_label('目标宿主', exact=True)).to_be_disabled()
    expect(page.get_by_label('目标宿主', exact=True)).to_have_value('D2P(default)')
    expect(page.get_by_text('序列只在本机计算与保存', exact=True)).to_be_visible()
    page.get_by_label('任务名称').fill('本地引擎验证')
    page.locator('#sequence-input').fill('>first\nMEFEFK\n>second\nMK*')
    page.get_by_role('button', name='EcoRI', exact=True).click()
    page.get_by_role('button', name='保存预设', exact=True).click()
    modal = page.get_by_role('dialog')
    modal.get_by_label('预设名称').fill('本地 EcoRI')
    modal.get_by_role('button', name='保存预设', exact=True).click()
    page.get_by_role('button', name='开始优化', exact=False).click()
    page.get_by_role('button', name='补充末尾 * 并开始优化', exact=True).click()
    expect(page.locator('.job-progress')).to_contain_text('2 条通过校验', timeout=15000)
    expect(page.locator('.job-engine-summary')).to_contain_text('本地实验')
    expect(page.locator('.job-engine-summary')).to_contain_text('在本机处理，未上传序列')
    jobs = page.request.get(URL + '/api/jobs').json()
    job_id = jobs[0]['id']
    job = page.request.get(URL + '/api/jobs/' + job_id).json()
    assert job['engine'] == 'local' and job['engine_info']['network_used'] is False
    assert all(record['response']['network_used'] is False for record in job['records'])
    page.get_by_role('tab', name='校验详情', exact=True).click()
    expect(page.get_by_text('d2p-local-experimental-1', exact=True)).to_be_visible()
    expect(page.get_by_text(job['engine_info']['mapping_sha256'], exact=True)).to_be_visible()
    expect(page.get_by_text('已按提交时的选择补充末尾终止符 *', exact=True)).to_be_visible()
    with page.expect_download() as download:
        page.get_by_role('button', name='导出 Excel', exact=True).click()
    output = OUT / 'local-engine.xlsx'
    download.value.save_as(str(output))
    workbook = load_workbook(output)
    assert workbook.sheetnames == ['基因合成']
    assert list(workbook.active.values)[0] == ('序号', '合成名', '核酸序列')
    assert workbook.active['C3'].value == 'ATGAAGTAA'
    assert all(cell.comment is None for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    with page.expect_download() as download:
        page.get_by_role('button', name='记录包', exact=True).click()
    bundle = OUT / 'local-engine-records.zip'
    download.value.save_as(str(bundle))
    with zipfile.ZipFile(bundle) as archive:
        saved = json.loads(archive.read('record.json'))
        assert saved['engine_info']['version'] == 'd2p-local-experimental-1'
        assert saved['records'][0]['attempts'][0]['engine'] == 'local'
    page.screenshot(path=str(OUT / 'local-engine-desktop.png'), full_page=True)
    page.reload(wait_until='networkidle')
    expect(page.get_by_role('radio', name='本地实验', exact=True)).to_be_checked()
    expect(page.locator('.job-progress')).to_contain_text('2 条通过校验')
    page.get_by_role('button', name='历史记录', exact=True).click()
    expect(page.locator('.history-row').filter(has_text='本地引擎验证')).to_contain_text('本地实验')
    page.get_by_role('button', name='参数预设', exact=True).click()
    expect(page.locator('.preset-card')).to_contain_text('本地实验')
    page.get_by_role('button', name='应用到新任务', exact=False).click()
    expect(page.get_by_role('radio', name='本地实验', exact=True)).to_be_checked()
    page.get_by_role('radio', name='ProteinFactory 远端', exact=True).check()
    expect(page.get_by_label('目标宿主', exact=True)).to_be_enabled()
    expect(page.get_by_text('序列将发送至 ProteinFactory 进行优化', exact=True)).to_be_visible()
    page.get_by_role('radio', name='本地实验', exact=True).check()
    page.get_by_role('button', name='新建优化', exact=True).click()
    page.locator('#sequence-input').fill('MMMM*')
    page.get_by_label('任务名称').fill('本地无解测试')
    page.get_by_role('button', name='EcoRI', exact=True).click()
    page.get_by_label('自定义避免位点', exact=True).fill('ATG')
    page.get_by_role('button', name='添加自定义位点', exact=True).click()
    page.get_by_role('button', name='开始优化', exact=False).click()
    expect(page.locator('.failure-state')).to_contain_text('无同义编码', timeout=10000)
    expect(page.locator('.job-engine-summary')).to_contain_text('本地实验')
    page.get_by_role('button', name='重试未完成项', exact=True).first.click()
    expect(page.locator('.failure-state')).to_contain_text('无同义编码', timeout=10000)
    assert page.request.get(URL + '/test/network-attempts').json()['attempts'] == []
    page.get_by_role('button', name='复制此任务参数与序列', exact=True).click()
    expect(page.get_by_role('radio', name='本地实验', exact=True)).to_be_checked()
    # Long results render in bounded chunks; full DNA remains in API/export data.
    page.get_by_role('button', name='新建优化', exact=True).click()
    page.locator('.motif-values button').click()
    page.locator('#sequence-input').fill('M' + 'K' * 1998 + '*')
    page.get_by_label('任务名称').fill('长序列渲染验证')
    page.get_by_role('button', name='开始优化', exact=False).click()
    expect(page.locator('.job-progress')).to_contain_text('1 条通过校验', timeout=15000)
    expect(page.locator('.dna-view .codon')).to_have_count(400)
    page.get_by_role('button', name='显示更多', exact=False).click()
    expect(page.locator('.dna-view .codon')).to_have_count(800)
    latest = page.request.get(URL + '/api/jobs').json()[0]
    full = page.request.get(URL + '/api/jobs/' + latest['id']).json()
    compact = page.request.get(URL + '/api/jobs/' + latest['id'] + '?compact=true').json()
    assert len(full['records'][0]['dna']) == len(compact['records'][0]['dna']) == 6000
    assert 'seq' not in compact['records'][0]['response']
    page.set_viewport_size({'width': 390, 'height': 844})
    expect(page.get_by_role('radio', name='本地实验', exact=True)).to_be_visible()
    assert not page.evaluate('document.documentElement.scrollWidth > innerWidth'), 'Mobile overflow'
    page.screenshot(path=str(OUT / 'local-engine-mobile.png'), full_page=True)
    assert not errors, errors
    assert not external_requests, external_requests
    report = {'status': 'passed', 'server_external_network_attempts': [], 'browser_external_requests': external_requests, 'browser_errors': errors,
              'checks': ['remote default', 'host compatibility', 'local compute with no HTTP calls', 'local constraints', 'terminal stop addition',
                         'engine/version/hash trace', 'clean XLSX', 'ZIP engine provenance', 'draft and refresh recovery', 'history engine label',
                         'engine preset', 'explicit remote switch', 'no fallback on infeasibility or retry', 'duplicate engine retention', 'bounded long DNA rendering with complete data', 'mobile layout']}
    (OUT / 'local-engine-e2e.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print('PASS: local UI, zero external calls, engine trace, stop prompt, exports, presets, failure/retry and mobile')
    browser.close()
