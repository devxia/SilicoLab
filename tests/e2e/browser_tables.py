"""Table mapping, example import, Excel exports and stop choices against the isolated UI fixture."""
import io
import json
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs' / 'ui'
OUT.mkdir(parents=True, exist_ok=True)
URL = 'http://127.0.0.1:8766'

with tempfile.TemporaryDirectory() as directory, sync_playwright() as p:
    browser = p.chromium.launch(headless=True, channel='chrome')
    page = browser.new_page(viewport={'width': 1440, 'height': 1060})
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto(URL, wait_until='networkidle')
    page.get_by_role('tab', name='导入表格', exact=True).click()
    original = Path(directory) / 'synthetic-input.xlsx'
    sample = Workbook()
    sample.active.title = '基因合成'
    sample.active.append(['序号', '合成名', '长度', '序列'])
    for index in range(1, 81):
        sample.active.append([index, f'synthetic-{index}', 135, 'M' + 'K' * 134 + '*'])
    sample.save(original)
    sample.close()
    page.get_by_label('选择表格文件').set_input_files(str(original))
    modal = page.get_by_role('dialog', name='导入表格 · 选择数据列')
    expect(modal.get_by_label('名称列', exact=True)).to_have_value('1')
    expect(modal.get_by_label('蛋白序列列', exact=True)).to_have_value('3')
    expect(modal.get_by_label('表头', exact=True)).to_have_value('0')
    expect(modal.get_by_text('80 条有效序列', exact=True)).to_be_visible()
    expect(modal.get_by_text('0 条缺少末尾 *', exact=True)).to_be_visible()
    page.screenshot(path=str(OUT / 'table-example-preview.png'), full_page=True)
    modal.get_by_role('button', name='导入 80 条序列', exact=True).click()
    expect(page.get_by_text('80 条序列已通过输入校验', exact=True)).to_be_visible()
    assert page.request.get(URL + '/api/jobs').json() == [], 'Import must not submit sequences'
    assert page.locator('.stop-reminder').count() == 0
    workbook = Workbook()
    workbook.active.title = '封面'
    workbook.active.append(['This is not the data sheet'])
    data = workbook.create_sheet('无表头数据')
    for row in [['备注', 'MK', 1, 'alpha-1'], ['备注', 'MK*', 2, 'beta-2'], [None], ['备注', 'MX', 3, 'bad-3']]:
        data.append(row)
    path = Path(directory) / 'headerless.xlsx'
    workbook.save(path)
    workbook.close()
    page.get_by_label('选择表格文件').set_input_files(str(path))
    modal = page.get_by_role('dialog', name='导入表格 · 选择数据列')
    expect(modal.get_by_label('工作表', exact=True)).to_have_value('1')
    expect(modal.get_by_label('表头', exact=True)).to_have_value('none')
    modal.get_by_label('名称列', exact=True).select_option('3')
    modal.get_by_label('蛋白序列列', exact=True).select_option('1')
    expect(modal.get_by_text('2 条有效序列', exact=True)).to_be_visible()
    expect(modal.get_by_text('1 行需要检查', exact=True)).to_be_visible()
    expect(modal.get_by_role('button', name='导入 3 条序列', exact=True)).to_be_disabled()
    modal.get_by_role('checkbox', name='仅导入 2 条有效序列，跳过上述 1 行', exact=True).check()
    page.set_viewport_size({'width': 390, 'height': 844})
    assert not page.evaluate('document.documentElement.scrollWidth > innerWidth'), 'Table preview causes mobile overflow'
    page.screenshot(path=str(OUT / 'table-mapping-mobile.png'), full_page=True)
    modal.get_by_role('button', name='导入 2 条序列', exact=True).click()
    page.set_viewport_size({'width': 1440, 'height': 1060})
    expect(page.get_by_text('1 条序列缺少末尾终止符 *', exact=True)).to_be_visible()
    page.get_by_label('任务名称').fill('表格导入与终止符验证')
    page.get_by_role('button', name='开始优化', exact=False).click()
    dialog = page.get_by_role('dialog', name='是否补充末尾终止符？')
    expect(dialog).to_be_visible()
    dialog.get_by_role('button', name='返回编辑', exact=True).click()
    assert page.request.get(URL + '/api/jobs').json() == []
    page.get_by_role('button', name='开始优化', exact=False).click()
    with page.expect_response(lambda response: response.url.endswith('/api/jobs') and response.request.method == 'POST') as created:
        page.get_by_role('button', name='补充末尾 * 并开始优化', exact=True).click()
    job_id = created.value.json()['id']
    expect(page.locator('.job-progress')).to_contain_text('2 条通过校验', timeout=15000)
    detail = page.request.get(URL + '/api/jobs/' + job_id).json()
    assert [record['sequence'] for record in detail['records']] == ['MK*', 'MK*']
    assert [record['stop_added'] for record in detail['records']] == [True, False]
    page.get_by_role('tab', name='校验详情', exact=True).click()
    expect(page.get_by_text('已按提交时的选择补充末尾终止符 *', exact=True)).to_be_visible()
    with page.expect_download() as download:
        page.get_by_role('button', name='导出 Excel', exact=True).click()
    exported = OUT / 'table-export.xlsx'
    download.value.save_as(str(exported))
    book = load_workbook(exported)
    assert book.sheetnames == ['基因合成']
    assert list(book['基因合成'].values) == [('序号', '合成名', '核酸序列'), (1, 'alpha-1', 'ATGAAGTAA'), (2, 'beta-2', 'ATGAAGTAA')]
    assert all(cell.comment is None for sheet in book for row in sheet for cell in row)
    assert book['基因合成']['A1'].font.name == '等线'
    assert book['基因合成'].column_dimensions['C'].width == 238.83203125
    book.close()
    page.screenshot(path=str(OUT / 'table-results.png'), full_page=True)
    page.get_by_role('button', name='新建优化', exact=True).click()
    page.get_by_role('tab', name='粘贴序列', exact=True).click()
    page.locator('#sequence-input').fill('MK')
    page.get_by_role('button', name='开始优化', exact=False).click()
    page.get_by_role('button', name='保持原序列，直接优化', exact=True).click()
    expect(page.locator('.job-progress')).to_contain_text('1 条通过校验', timeout=15000)
    page.get_by_role('tab', name='校验详情', exact=True).click()
    expect(page.get_by_text('输入未带终止符，未额外添加终止密码子', exact=True)).to_be_visible()
    page.get_by_role('button', name='新建优化', exact=True).click()
    page.locator('#sequence-input').fill('>first\nMK\n>second\nMK*')
    page.get_by_role('button', name='一键补齐', exact=True).click()
    expect(page.locator('#sequence-input')).to_have_value('>first\nMK*\n\n>second\nMK*')
    assert page.locator('.stop-reminder').count() == 0
    assert not errors, errors
    (OUT / 'table-e2e-report.json').write_text(json.dumps({'status': 'passed', 'upstream': 'isolated fixture only',
        'checks': ['80 synthetic sequences imported locally without submission', 'multi-sheet selection', 'headerless mapping', 'manual column selection', 'blank rows', 'explicit invalid-row skipping', 'mobile mapping', 'stop prompt cancel/add/preserve', 'Excel template headers and data', 'single-sheet export without cell comments', 'one-click stop completion'], 'browser_errors': errors}, ensure_ascii=False, indent=2))
    print('PASS: example import, arbitrary columns, headerless sheets, stop choices, matching XLSX export and mobile')
    print('Workbook:', exported)
    browser.close()
