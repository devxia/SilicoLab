"""Explicit science upload, verified exports and unknown→explicit retry in real UI."""
import base64
import io
import json
from pathlib import Path
import tempfile
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

URL = 'http://127.0.0.1:8767'
OUT = Path('outputs/ui/remote-science')
OUT.mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory(prefix='silicolab-remote-ui-project-') as directory, sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel='chrome', headless=True)
    page = browser.new_page(viewport={'width': 1480, 'height': 1000})
    errors, external = [], []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.on('request', lambda request: external.append(request.url) if not request.url.startswith(URL) else None)
    project = page.request.post(URL + '/api/workbench/projects', data={'path': directory, 'name': '远端科学专项'}).json()
    url = URL + '/api/workbench/projects/' + project['id']
    chat = page.request.post(url + '/conversations', data={'title': '人工数据与远端未知'}).json()['id']
    imported = page.request.post(url + '/files', data={'name': 'artificial.fa', 'contentBase64': base64.b64encode(b'>artificial\nMK*\n').decode()}).json()
    fixed = page.request.post(url + '/input-versions', data={'conversationId': chat, 'fileIds': [imported['id']]}).json()
    page.goto(URL, wait_until='networkidle')
    page.get_by_role('button', name='打开项目 远端科学专项', exact=True).click()
    page.get_by_role('button', name='查看任务', exact=True).click()
    form = page.get_by_role('region', name='ProteinFactory 远端专项任务')
    expect(form).to_contain_text('这不是模型 provider 授权')
    expect(form).to_contain_text('不会自动回退到本地')
    expect(form).to_contain_text('pD2P_1.08t')
    form.get_by_label('固定输入版本').select_option(fixed['id'])
    expect(form.get_by_role('button', name='运行 ProteinFactory 远端专项', exact=True)).to_be_disabled()
    assert page.request.get(URL + '/__test/science-requests').json()['requests'] == []
    form.get_by_label('我明确使用 ProteinFactory 远端专项，并授权上传所选序列').check()
    form.get_by_role('button', name='运行 ProteinFactory 远端专项', exact=True).click()
    transcript = page.locator('.wb-transcript')
    expect(transcript).to_contain_text('artificial · 已校验', timeout=15000)
    expect(transcript).to_contain_text('实际计算来源：ProteinFactory 远端')
    expect(transcript).to_contain_text('网络请求：1')
    with page.expect_download() as download:
        transcript.get_by_role('link', name='下载合成表 Excel', exact=True).click()
    path = OUT / 'synthesis.xlsx'
    download.value.save_as(str(path))
    workbook = load_workbook(io.BytesIO(path.read_bytes()))
    assert workbook.sheetnames == ['基因合成']
    assert list(workbook.active.values)[1] == (1, 'artificial', 'ATGAAGTAA')
    assert all(cell.comment is None for row in workbook.active for cell in row)
    workbook.close()
    form.get_by_label('正反向避免位点').fill('GAATTC')
    form.get_by_role('button', name='运行 ProteinFactory 远端专项', exact=True).click()
    unknown = transcript.locator('.wb-task-card').last
    expect(unknown).to_contain_text('结果未知，等待核对', timeout=15000)
    expect(unknown).to_contain_text('结果未知，不自动重试')
    page.get_by_role('button', name='查看结果', exact=True).click()
    results = page.get_by_label('真实科研结果')
    expect(results).to_contain_text('结果未知，等待核对')
    expect(results).to_contain_text('0 条失败')
    expect(results).to_contain_text('1 条结果未知')
    expect(results.get_by_role('region', name='序列结果表')).to_contain_text('结果未知')
    page.screenshot(path=str(OUT / 'unknown-results.png'), full_page=True)
    page.get_by_role('button', name='查看任务', exact=True).click()
    unknown.get_by_role('button', name='核对远端请求记录', exact=True).click()
    expect(unknown).to_contain_text('没有已验证的查询接口')
    expect(unknown).to_contain_text('可能重复计算')
    assert len(page.request.get(URL + '/__test/science-requests').json()['requests']) == 2
    page.screenshot(path=str(OUT / 'unknown-decision.png'), full_page=True)
    unknown.get_by_role('button', name='明确重试远端请求（可能重复计算）', exact=True).click()
    expect(unknown).to_contain_text('artificial · 已校验', timeout=15000)
    requests = page.request.get(URL + '/__test/science-requests').json()['requests']
    assert len(requests) == 3
    tasks = page.request.get(url).json()['tasks']
    assert tasks[0]['toolExecutions'][0]['result']['records'][0]['dna'] == 'ATGAAGTAA'
    assert len(tasks[1]['toolExecutions']) == 2
    assert tasks[1]['toolExecutions'][0]['status'] == 'unknown'
    assert tasks[1]['toolExecutions'][1]['status'] == 'completed'
    assert page.request.get(URL + '/api/jobs').json() == []
    page.screenshot(path=str(OUT / 'retried-result.png'), full_page=True)
    assert not errors, errors
    assert not external, external
    (OUT / 'result.json').write_text(json.dumps({'result': 'passed', 'science_requests': 3, 'source': 'artificial loopback endpoint', 'unknown_preserved': True, 'explicit_retry_only': True, 'errors': errors, 'external_browser_requests': external}, indent=2))
    browser.close()
