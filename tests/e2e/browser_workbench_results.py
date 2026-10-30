"""Result workspace shows real partial computation, complete copy and matching exports."""
import base64
import io
import json
from pathlib import Path
import tempfile
import time
import zipfile
from openpyxl import load_workbook
from fixtures.legacy_scientific_project import create_legacy_project
from playwright.sync_api import expect, sync_playwright

URL = 'http://127.0.0.1:8767'
OUT = Path('outputs/ui/scientific-results')
OUT.mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory(prefix='silicolab-results-ui-') as directory, sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel='chrome', headless=True)
    context = browser.new_context(viewport={'width': 1490, 'height': 1000}, permissions=['clipboard-read', 'clipboard-write'])
    page = context.new_page()
    errors, external = [], []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.on('request', lambda request: external.append(request.url) if not request.url.startswith(URL) else None)
    # Arrange through the same real application contract, not injected UI data.
    project = page.request.post(URL + '/api/workbench/projects', data={'path': directory, 'name': '部分结果可核对'}).json()
    url = URL + '/api/workbench/projects/' + project['id']
    chat = page.request.post(url + '/conversations', data={'title': '三条序列部分失败'}).json()['id']
    content = b'>good-1\nK*\n>bad\nM*\n>good-2\n' + b'K' * 3000 + b'*\n'
    long_dna = 'AAG' * 3000 + 'TAA'
    imported = page.request.post(url + '/files', data={'name': 'mixed.fasta', 'contentBase64': base64.b64encode(content).decode()}).json()
    version = page.request.post(url + '/input-versions', data={'conversationId': chat, 'fileIds': [imported['id']]}).json()
    task = page.request.post(url + '/tasks', data={'conversationId': chat, 'requestId': 'result-case', 'goal': 'D2P 专项本地实验，核对部分失败', 'authorizedTools': ['d2p_local'], 'workflow': 'd2p_local', 'specialtyAcknowledged': True, 'inputVersionId': version['id'], 'parameters': {'stopPolicy': 'preserve', 'avoid': 'ATG'}}).json()
    deadline = time.monotonic() + 12
    while time.monotonic() < deadline:
        state = page.request.get(url + '/conversations/' + chat).json()
        if state['tasks'][0]['status'] == 'answered':
            break
        time.sleep(0.03)
    else:
        raise AssertionError(state)
    execution = state['tasks'][0]['toolExecutions'][0]
    assert execution['result']['status'] == 'partial'
    before_requests = state['modelRequests']
    page.goto(URL, wait_until='networkidle')
    page.get_by_role('button', name='打开项目 部分结果可核对', exact=True).click()
    page.get_by_role('button', name='查看结果', exact=True).click()
    results = page.get_by_label('真实科研结果', exact=True)
    expect(results).to_contain_text('部分完成')
    expect(results).to_contain_text('2 条通过　1 条失败')
    expect(results.get_by_role('row')).to_have_count(4)
    expect(results.get_by_role('row').filter(has_text='bad')).to_contain_text('无结果')
    expect(results.get_by_role('row').filter(has_text='bad')).to_contain_text('失败')
    results.get_by_role('button', name='检视序列 good-1', exact=True).click()
    results.get_by_role('button', name='复制完整 DNA', exact=True).click()
    expect(results.get_by_role('status')).to_contain_text('6 bp')
    assert page.evaluate('navigator.clipboard.readText()') == 'AAGTAA'
    sequence_tab = results.get_by_role('tab', name='序列表', exact=True)
    sequence_tab.focus()
    sequence_tab.press('ArrowRight')
    dna_tab = results.get_by_role('tab', name='DNA 与密码子', exact=True)
    expect(dna_tab).to_be_focused()
    expect(dna_tab).to_have_attribute('aria-selected', 'true')
    assert results.get_by_role('tabpanel').get_attribute('id') == dna_tab.get_attribute('aria-controls')
    dna_tab.press('End')
    provenance_tab = results.get_by_role('tab', name='校验与来源', exact=True)
    expect(provenance_tab).to_be_focused()
    expect(provenance_tab).to_have_attribute('aria-selected', 'true')
    provenance_tab.press('Home')
    expect(sequence_tab).to_be_focused()
    expect(sequence_tab).to_have_attribute('aria-selected', 'true')
    results.get_by_role('tab', name='DNA 与密码子', exact=True).click()
    expect(results.get_by_label('密码子对照')).to_contain_text('AAG')
    expect(results.get_by_label('完整 DNA 序列')).to_have_text('AAGTAA')
    results.get_by_role('tab', name='校验与来源', exact=True).click()
    expect(results).to_contain_text('翻译一致性：通过')
    expect(results).to_contain_text('终止符策略：通过')
    expect(results).to_contain_text(execution['id'])
    expect(results).to_contain_text(imported['sha256'])
    results.get_by_role('button', name='检视序列 bad', exact=True).click()
    expect(results.get_by_role('alert')).to_be_visible()
    expect(results).to_contain_text('没有通过校验的 DNA')
    expect(results).to_contain_text('存在末尾终止符')
    expect(results).not_to_contain_text('未带末尾终止符')
    expect(results.get_by_role('button', name='复制完整 DNA', exact=True)).to_have_count(0)
    results.get_by_role('button', name='检视序列 good-2', exact=True).click()
    results.get_by_role('button', name='复制完整 DNA', exact=True).click()
    expect(results.get_by_role('status')).to_contain_text('9003 bp')
    assert page.evaluate('navigator.clipboard.readText()') == long_dna
    with page.expect_download() as download:
        results.get_by_role('link', name='单条 FASTA', exact=True).click()
    single = OUT / 'single.fasta'
    download.value.save_as(str(single))
    assert single.read_text() == '>good-2\n' + long_dna + '\n'
    with page.expect_download() as download:
        results.get_by_role('link', name='合成表 Excel', exact=True).click()
    excel = OUT / 'synthesis.xlsx'
    download.value.save_as(str(excel))
    workbook = load_workbook(excel)
    assert workbook.sheetnames == ['基因合成']
    assert list(workbook.active.values)[1:] == [(1, 'good-1', 'AAGTAA'), (2, 'good-2', long_dna)]
    assert all(cell.comment is None for row in workbook.active for cell in row)
    workbook.close()
    with page.expect_download() as download:
        results.get_by_role('link', name='完整记录包', exact=True).click()
    bundle = OUT / 'records.zip'
    download.value.save_as(str(bundle))
    with zipfile.ZipFile(bundle) as archive:
        assert json.loads(archive.read('record.json')) == execution['result']
        assert archive.read('inputs/001.fasta') == content
        assert archive.read('synthesis.xlsx') == excel.read_bytes()
    page.screenshot(path=str(OUT / 'desktop.png'), full_page=True)
    page.set_viewport_size({'width': 620, 'height': 850})
    expect(results.get_by_role('region', name='序列结果表')).to_be_visible()
    results.get_by_role('region', name='序列结果表').evaluate('(el) => el.scrollLeft = el.scrollWidth')
    assert results.get_by_role('region', name='序列结果表').evaluate('(el) => el.scrollLeft') > 0
    page.screenshot(path=str(OUT / 'narrow.png'), full_page=True)
    page.reload(wait_until='domcontentloaded')
    page.get_by_role('button', name='打开项目 部分结果可核对', exact=True).click()
    page.get_by_role('button', name='收起导航', exact=True).click()
    page.get_by_role('button', name='查看结果', exact=True).click()
    expect(page.get_by_label('真实科研结果', exact=True)).to_contain_text('部分完成')
    after = page.request.get(url + '/conversations/' + chat).json()
    assert after['modelRequests'] == before_requests
    assert after['tasks'][0]['toolExecutions'] == state['tasks'][0]['toolExecutions']
    legacy_path = Path(directory) / 'previous-project'
    legacy_id, legacy_chat, legacy_state = create_legacy_project(legacy_path)
    legacy_execution = legacy_state['tasks'][0]['toolExecutions'][0]
    reopened = page.request.post(URL + '/api/workbench/projects/open', data={'path': str(legacy_path)})
    assert reopened.ok, reopened.text()
    page.reload(wait_until='domcontentloaded')
    page.get_by_role('button', name='打开项目 旧格式科学结果', exact=True).click()
    page.get_by_role('button', name='收起导航', exact=True).click()
    page.get_by_role('button', name='查看结果', exact=True).click()
    legacy_results = page.get_by_label('真实科研结果', exact=True)
    expect(legacy_results).to_contain_text('部分完成')
    legacy_results.get_by_role('tab', name='校验与来源', exact=True).click()
    expect(legacy_results).to_contain_text(legacy_execution['id'])
    expect(legacy_results).to_contain_text(legacy_execution['callId'])
    legacy_results.get_by_role('button', name='检视序列 legacy-failed', exact=True).click()
    expect(legacy_results).to_contain_text('存在末尾终止符')
    expect(legacy_results).not_to_contain_text('未带末尾终止符')
    legacy_after = page.request.get(f'{URL}/api/workbench/projects/{legacy_id}/conversations/{legacy_chat}').json()
    assert legacy_after['tasks'] == legacy_state['tasks']
    assert legacy_after['modelRequests'] == legacy_state['modelRequests'] == 2
    assert not errors, errors
    assert not external, external
    (OUT / 'result.json').write_text(json.dumps({'result': 'passed', 'passed': 2, 'failed': 1, 'model_requests_unchanged': before_requests, 'legacy_project_reopened_without_rerun': True, 'keyboard_tabs': 'passed', 'external': external, 'errors': errors}, indent=2))
    browser.close()
