"""Explicit local specialty via workbench, actual validation and Excel/FASTA downloads."""
import io
import json
from pathlib import Path
import tempfile
from playwright.sync_api import expect, sync_playwright
from openpyxl import load_workbook

URL = 'http://127.0.0.1:8767'
OUT = Path('outputs/ui/local-science')
OUT.mkdir(parents=True, exist_ok=True)
with tempfile.TemporaryDirectory(prefix='silicolab-local-science-ui-') as directory, sync_playwright() as playwright:
    browser = playwright.chromium.launch(channel='chrome', headless=True)
    page = browser.new_page(viewport={'width': 1480, 'height': 1000})
    errors, external = [], []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.on('request', lambda request: external.append(request.url) if not request.url.startswith(URL) else None)
    page.goto(URL, wait_until='networkidle')
    page.get_by_role('button', name='创建项目', exact=True).first.click()
    dialog = page.get_by_role('dialog')
    dialog.get_by_label('项目名称', exact=True).fill('本地专项闭环')
    dialog.get_by_label('项目目录', exact=True).fill(directory)
    dialog.get_by_role('button', name='确认', exact=True).click()
    expect(dialog).not_to_be_visible()
    page.get_by_role('button', name='新建对话', exact=True).first.click()
    dialog = page.get_by_role('dialog')
    dialog.get_by_label('对话名称', exact=True).fill('校验与合成表')
    dialog.get_by_role('button', name='确认', exact=True).click()
    expect(dialog).not_to_be_visible()
    page.get_by_role('button', name='查看文件', exact=True).click()
    page.get_by_label('导入项目文件').set_input_files({'name': 'proteins.fasta', 'mimeType': 'text/plain', 'buffer': b'>first\nMK*\n>second\nMEFEFK*\n'})
    expect(page.locator('.wb-files li')).to_have_count(1)
    page.locator('.wb-files input[type=checkbox]').check()
    page.get_by_role('button', name='固定任务输入', exact=True).click()
    expect(page.get_by_role('region', name='固定输入详情')).to_be_visible()
    page.get_by_role('button', name='D2P 专项本地实验', exact=True).click()
    form = page.get_by_role('region', name='D2P 专项本地实验任务')
    expect(form).to_contain_text('实际计算来源：本地 Python 实验引擎')
    expect(form).to_contain_text('尚无载体校验')
    options = form.get_by_label('固定输入版本').locator('option').all()
    form.get_by_label('固定输入版本').select_option(options[1].get_attribute('value'))
    form.get_by_label('正反向避免位点').fill('GAATTC')
    form.get_by_label('我明确使用 D2P 专项的本地实验，不使用远端').check()
    form.get_by_role('button', name='运行 D2P 专项本地实验', exact=True).click()
    transcript = page.locator('.wb-transcript')
    expect(transcript).to_contain_text('first · 已校验', timeout=15000)
    expect(transcript).to_contain_text('网络请求：0')
    expect(transcript).to_contain_text('d2p-local-experimental-1')
    with page.expect_download() as download:
        transcript.get_by_role('link', name='下载合成表 Excel', exact=True).click()
    path = OUT / 'synthesis.xlsx'
    download.value.save_as(str(path))
    workbook = load_workbook(io.BytesIO(path.read_bytes()))
    assert workbook.sheetnames == ['基因合成']
    assert list(workbook.active.values)[0] == ('序号', '合成名', '核酸序列')
    assert workbook.active['C2'].value == 'ATGAAGTAA'
    assert all(cell.comment is None for row in workbook.active for cell in row)
    workbook.close()
    with page.expect_download() as download:
        transcript.get_by_role('link', name='下载 DNA FASTA', exact=True).click()
    fasta = OUT / 'synthesis.fasta'
    download.value.save_as(str(fasta))
    assert '>first\nATGAAGTAA' in fasta.read_text()
    page.screenshot(path=str(OUT / 'result.png'), full_page=True)
    identity = page.request.get(URL + '/api/workbench/projects').json()[0]['id']
    task = page.request.get(f'{URL}/api/workbench/projects/{identity}').json()['tasks'][0]
    assert task['toolExecutions'][0]['result']['network_guard_self_test'] == 'blocked'
    assert task['toolExecutions'][0]['result']['network_attempts'] == 0
    assert page.request.get(URL + '/api/jobs').json() == []
    assert not errors, errors
    assert not external, external
    (OUT / 'result.json').write_text(json.dumps({'result': 'passed', 'records': 2, 'network_attempts': 0, 'external': external, 'errors': errors}, indent=2))
    browser.close()
