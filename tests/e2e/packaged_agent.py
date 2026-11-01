"""Installed ZIP agent verification: real packaged runtimes, artificial input only."""
import argparse
from contextlib import ExitStack
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import time
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parent/'fixtures'))
sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'backend/tests'))
from packaged_app import PackagedApp,wait,windows,hidden_builder
from test_model_configuration import FakeModelEndpoint
from test_mcp_resources import SERVER
from test_restricted_analysis import SCRIPT


def verify(args):
    archive=args.archive.resolve();archive_sha=hashlib.sha256(archive.read_bytes()).hexdigest();report={'archiveSha256':archive_sha,'checks':{},'artificialDataOnly':True,'realProviderEvidence':'docs/provider-compatibility.md; no extra real calls'}
    with tempfile.TemporaryDirectory(prefix='silicolab-installed-验收 space-') as directory,ExitStack() as cleanup:
        root=Path(directory);subprocess.run(['/usr/bin/ditto','-x','-k',str(archive),str(root)],check=True)
        app=root/'SilicoLab.app';assert app.is_dir()
        subprocess.run(['/usr/bin/codesign','--verify','--deep','--strict',str(app)],check=True,capture_output=True)
        resources=app/'Contents/Resources'
        assert subprocess.check_output([str(resources/'runtime/bin/node'),'--version'],text=True).strip()=='v24.16.0'
        assert (resources/'runtime/node_modules/@earendil-works/pi-durable/package.json').is_file()
        assert (resources/'analysis/PythonBoundaryProbe.app').is_dir() and (resources/'notifications/SilicoLabNotifications.app').is_dir()
        assert (resources/'frontend/dist/fonts/DM-Sans.ttf').is_file() and (resources/'frontend/dist/fonts/OFL.txt').is_file()
        report['checks']['bundledRuntimeResources']=True
        builder=args.builder_root.resolve()
        cleanup.enter_context(hidden_builder(builder))
        owner=PackagedApp(app,root/'app-data');configs=[]
        try:
            gui=owner.launch();report['startupSeconds']=round(owner.startup_seconds,3)
            original_service=owner.info['serviceId'];report['checks']['realCocoaWindow']=bool(windows(gui.pid))
            assert owner.api('/projects')==[] and owner.api('/model-configs')==[]
            project_path=root/'artificial-project';project_path.mkdir()
            skill=project_path/'.agents/skills/artificial';skill.mkdir(parents=True);(skill/'SKILL.md').write_text('---\nname: artificial\ndescription: Artificial packaged resource\n---\nOnly artificial knowledge.\n')
            (project_path/'server.py').write_text(SERVER);(project_path/'.mcp.json').write_text(json.dumps({'mcpServers':{'artificial':{'type':'stdio','command':'silicolab-restricted-python','args':['server.py']}}}))
            project=owner.api('/projects',{'path':str(project_path),'name':'安装包人工项目'});url='/projects/'+project['id'];original_node=project['runtimePid']
            chat=owner.api(url+'/conversations',{'title':'原科研对话'})['id']
            def fixed(name,content):
                source=owner.api(url+'/files',{'name':name,'contentBase64':base64.b64encode(content).decode()})
                return owner.api(url+'/input-versions',{'conversationId':chat,'fileIds':[source['id']]})
            fasta=fixed('artificial.fa',b'>synthetic\nMK*\n')
            task=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'packaged-local','goal':'明确D2P本地专项','workflow':'d2p_local','specialtyAcknowledged':True,'authorizedTools':['d2p_local'],'inputVersionId':fasta['id'],'parameters':{'stopPolicy':'preserve','avoid':''}})
            done=owner.task(url,task['id']);execution=done['toolExecutions'][0]
            assert execution['status']=='completed' and execution['result']['records'][0]['dna']=='ATGAAGTAA'
            assert execution['result']['network_attempts']==0
            data=owner.raw('/api/workbench'+url+'/tasks/'+task['id']+'/artifacts/'+execution['id']+'/synthesis.xlsx').content
            book=load_workbook(io.BytesIO(data));assert book.sheetnames==['基因合成'];assert list(book.active.values)==[('序号','合成名','核酸序列'),(1,'synthetic','ATGAAGTAA')]
            assert all(cell.comment is None for row in book.active for cell in row);book.close()
            report['checks']['localD2PAndExcel']=True
            csv=fixed('artificial.csv',b'value\n1\n2\n6\n')
            analysis=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'packaged-analysis','goal':'人工CSV摘要','workflow':'python_analysis','authorizedTools':['python_analysis'],'inputVersionId':csv['id']})
            result=owner.task(url,analysis['id'])['toolExecutions'][0]
            assert result['status']=='completed' and result['result']['summary']=={'count':3,'sum':9,'mean':3}
            assert result['result']['environment']['python']=='3.13.18';report['checks']['independentSandboxPython']=True
            task=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'skill','goal':'/skill artificial','authorizedTools':['synthetic_count']})
            assert owner.task(url,task['id'])['skillLoads'][0]['name']=='artificial';report['checks']['skills']=True
            mcp=owner.api(url+'/resources/mcp')['resources'][0];target=url+'/resources/mcp/'+mcp['id']
            owner.api(target+'/trust',{'fingerprint':mcp['fingerprint']});assert owner.api(target+'/connect',{'fingerprint':mcp['fingerprint']})['tools'][0]['name']=='echo'
            task=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'mcp','goal':'人工echo','workflow':'mcp_tool','authorizedTools':['mcp_call'],'mcpResourceId':mcp['id'],'mcpFingerprint':mcp['fingerprint'],'mcpTool':'echo','mcpArguments':{'text':'ABC'}})
            assert owner.task(url,task['id'])['toolExecutions'][0]['status']=='completed';report['checks']['mcpStdioSandbox']=True
            waiting=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'waiting','goal':'人工等待','authorizedTools':[]});decision=waiting['decisions'][0]
            assert waiting['status']=='needs_decision'
            body={'answer':'decline'};answer=url+'/tasks/'+waiting['id']+'/decisions/'+decision['id']
            assert owner.api(answer,body)==owner.api(answer,body);report['checks']['sameDecisionIdempotence']=True
            # Native close, task keeps running, duplicate app attaches to the same owner.
            with FakeModelEndpoint('openai-responses','artificial-close-barrier','hold') as barrier:
                config=owner.api('/model-configs',{'name':'关窗人工屏障','protocol':barrier.protocol,'baseUrl':barrier.url,'modelId':'artificial-model','apiKey':barrier.key});configs=[config]
                barrier_chat=owner.api(url+'/conversations',{'title':'关窗持续执行'})['id']
                owner.api(url+'/conversations/'+barrier_chat+'/model',{'configurationId':config['id']},'PUT')
                running=owner.api(url+'/tasks',{'conversationId':barrier_chat,'requestId':'native-background','goal':'ARTIFICIAL_BACKGROUND','authorizedTools':['synthetic_count']})
                assert barrier.entered.wait(5)
                close_start=time.monotonic();assert owner.command('--close-window').returncode==0;gui.wait(timeout=15);assert not windows(gui.pid)
                report['closeSeconds']=round(time.monotonic()-close_start,3)
                assert next(t for t in owner.api(url)['tasks'] if t['id']==running['id'])['status']=='running'
                assert not barrier.release.is_set();closed_at=time.time();barrier.release.set()
                completed=owner.task(url,running['id'])
                from datetime import datetime
                assert datetime.fromisoformat(completed['finishedAt'].replace('Z','+00:00')).timestamp()>closed_at
                assert len(completed['toolExecutions'])==1 and owner.api(url)['runtimePid']==original_node
                owner.api(url+'/conversations/'+barrier_chat+'/model',{'configurationId':None},'PUT');owner.api('/model-configs/'+config['id'],method='DELETE');configs=[]
            reopened=owner.launch();assert owner.info['serviceId']==original_service
            duplicate=owner.command('--isolated-resources');assert duplicate.returncode==0
            assert json.loads((owner.directory/'gui-session.json').read_text())['pid']==reopened.pid
            report['checks']['nativeCloseReopenSameOwner']=True
            # Protocol switch uses genuine app routes with artificial endpoints only.
            with FakeModelEndpoint('openai-responses','artificial-bundle-a','hold') as first,FakeModelEndpoint('anthropic-messages','artificial-bundle-b') as second:
                configs=[owner.api('/model-configs',{'name':'人工协议','protocol':m.protocol,'baseUrl':m.url,'modelId':'artificial-model','apiKey':m.key}) for m in (first,second)]
                switched=owner.api(url+'/conversations',{'title':'协议切换'})['id'];owner.api(url+'/conversations/'+switched+'/model',{'configurationId':configs[0]['id']},'PUT')
                task=owner.api(url+'/tasks',{'conversationId':switched,'requestId':'switch','goal':'ABC','authorizedTools':['synthetic_count']});assert first.entered.wait(5)
                owner.api(url+'/conversations/'+switched+'/model',{'configurationId':configs[1]['id']},'PUT');first.release.set();result=owner.task(url,task['id'])
                assert [r['protocol'] for r in result['modelRequests']]==['openai-responses','anthropic-messages'] and len(result['toolExecutions'])==1
                for config in configs:owner.api('/model-configs/'+config['id'],method='DELETE')
                configs=[]
            report['checks']['protocolSwitchArtificial']=True
            # Full quit with real sandbox script, queued task and pending decision.
            with FakeModelEndpoint('openai-responses','artificial-bundle-quit',tool_name='python_analysis',tool_arguments={'code':'import time;time.sleep(12)\n'+SCRIPT}) as model:
                config=owner.api('/model-configs',{'name':'人工退出','protocol':model.protocol,'baseUrl':model.url,'modelId':'artificial-model','apiKey':model.key});configs=[config]
                quit_chat=owner.api(url+'/conversations',{'title':'退出恢复'})['id'];owner.api(url+'/conversations/'+quit_chat+'/model',{'configurationId':config['id']},'PUT')
                source=owner.api(url+'/files',{'name':'quit.csv','contentBase64':base64.b64encode(b'value\n1\n2\n6\n').decode()});version=owner.api(url+'/input-versions',{'conversationId':quit_chat,'fileIds':[source['id']]})
                task=owner.api(url+'/tasks',{'conversationId':quit_chat,'requestId':'quit-script','goal':'人工退出分析','workflow':'python_analysis','authorizedTools':['python_analysis'],'inputVersionId':version['id']})
                queued=owner.api(url+'/tasks',{'conversationId':quit_chat,'requestId':'quit-queue','goal':'ABC','authorizedTools':['synthetic_count']})
                pending_decision=owner.api(url+'/tasks',{'conversationId':chat,'requestId':'wait-across-quit','goal':'人工待决定跨退出','authorizedTools':[]})
                active=wait(lambda:next((t for t in owner.api(url)['tasks'] if t['id']==task['id'] and t['toolExecutions'] and t['toolExecutions'][0].get('pythonPid')),None));pid=active['toolExecutions'][0]['pythonPid'];count=len(model.requests)
                owner.quit()
                for identity in (pid,original_node):
                    try:os.kill(identity,0)
                    except ProcessLookupError:pass
                    else:raise AssertionError('Managed script/node survived packaged full quit')
                owner.launch();opened=owner.api('/projects/open',{'path':str(project_path)})
                assert next(t for t in opened['tasks'] if t['id']==task['id'])['status']=='paused';assert len(model.requests)==count
                assert next(t for t in opened['tasks'] if t['id']==pending_decision['id'])['decisions']==pending_decision['decisions']
                owner.api(url+'/conversations/'+quit_chat+'/model',{'configurationId':None},'PUT');owner.api(url+'/tasks/'+queued['id']+'/resume',{})
                owner.task(url,queued['id']);assert next(t for t in owner.api(url)['tasks'] if t['id']==task['id'])['status']=='paused';assert len(model.requests)==count
                assert next(t for t in owner.api(url)['tasks'] if t['id']==pending_decision['id'])['status']=='needs_decision'
                owner.api('/model-configs/'+config['id'],method='DELETE');configs=[]
            report['checks']['fullQuitStopsScriptsExplicitRecovery']=True
            # Directory movement must be visible offline and explicitly verified before recovery.
            before_move=owner.api(url);moved=root/'moved-project';project_path.rename(moved)
            wait(lambda:owner.api(url).get('availability',{}).get('status')=='offline')
            restored=owner.api(url+'/relocate',{'path':str(moved)})
            assert restored['id']==project['id'] and restored['modelRequests']==before_move['modelRequests']
            assert next(t for t in restored['tasks'] if t['id']==task['id'])['status']=='paused'
            report['checks']['offlineRelocationNoAutoResume']=True
            # Consistent backup preserves UUID, completed artifacts and paused choices.
            archive_path=root/'project-backup.zip';owner.api(url+'/backup',{'path':str(archive_path)})
            assert archive_path.is_file();owner.api(url,method='DELETE');assert moved.exists()
            restore_path=root/'restored-project';restore_path.mkdir()
            archived=owner.api('/projects/restore',{'archivePath':str(archive_path),'path':str(restore_path)})
            assert archived['id']==project['id'] and len(archived['tasks'])==len(restored['tasks'])
            assert archived['modelRequests']==restored['modelRequests']
            assert next(t for t in archived['tasks'] if t['id']==analysis['id'])['toolExecutions'][0]['result']['summary']=={'count':3,'sum':9,'mean':3}
            report['checks']['backupRestoreAndNonDestructiveRemoval']=True
            # Resume the original interrupted Python task, not only a different queued task.
            frozen_executions={t['id']:t['toolExecutions'] for t in archived['tasks'] if t['status']=='answered'}
            old_attempts=next(t for t in archived['tasks'] if t['id']==task['id'])['toolExecutions']
            owner.api(url+'/tasks/'+task['id']+'/resume',{})
            resumed=owner.task(url,task['id'])
            assert resumed['toolExecutions'][:len(old_attempts)]==old_attempts
            assert resumed['toolExecutions'][-1]['result']['summary']=={'count':3,'sum':9,'mean':3}
            for saved in owner.api(url)['tasks']:
                if saved['id'] in frozen_executions:assert saved['toolExecutions']==frozen_executions[saved['id']]
            assert next(t for t in owner.api(url)['tasks'] if t['id']==pending_decision['id'])['status']=='needs_decision'
            report['checks']['originalInterruptedScriptExplicitlyResumed']=True
            # Pause after successful Python but during the next model request; the duplicate
            # tool call on resume must reuse exactly the completed execution/artifacts.
            with FakeModelEndpoint('openai-responses','artificial-reuse',mode=['normal','hold','duplicate_once','normal'],tool_name='python_analysis',tool_arguments={'code':SCRIPT}) as model:
                config=owner.api('/model-configs',{'name':'完成项复用','protocol':model.protocol,'baseUrl':model.url,'modelId':'artificial-model','apiKey':model.key});configs=[config]
                reuse_chat=owner.api(url+'/conversations',{'title':'完成Python暂停复用'})['id']
                owner.api(url+'/conversations/'+reuse_chat+'/model',{'configurationId':config['id']},'PUT')
                file=owner.api(url+'/files',{'name':'reuse.csv','contentBase64':base64.b64encode(b'value\n1\n2\n6\n').decode()});fixed=owner.api(url+'/input-versions',{'conversationId':reuse_chat,'fileIds':[file['id']]})
                reuse=owner.api(url+'/tasks',{'conversationId':reuse_chat,'requestId':'reuse-after-pause','goal':'人工受限分析','workflow':'python_analysis','authorizedTools':['python_analysis'],'inputVersionId':fixed['id']})
                assert model.entered.wait(15);before=next(t for t in owner.api(url)['tasks'] if t['id']==reuse['id'])
                assert len(before['toolExecutions'])==1 and before['toolExecutions'][0]['status']=='completed'
                owner.api(url+'/tasks/'+reuse['id']+'/pause',{});model.release.set()
                owner.api(url+'/tasks/'+reuse['id']+'/resume',{});after=owner.task(url,reuse['id'])
                assert after['toolExecutions']==before['toolExecutions']
                for artifact in after['toolExecutions'][0]['artifacts']:
                    data=owner.raw('/api/workbench'+url+'/tasks/'+reuse['id']+'/artifacts/'+after['toolExecutions'][0]['id']+'/'+artifact['name']).content
                    assert hashlib.sha256(data).hexdigest()==artifact['sha256']
                owner.api('/model-configs/'+config['id'],method='DELETE');configs=[]
            report['checks']['completedPythonReusedAfterPause']=True
            report['checks']['builderAbsentDuringVerification']=True
            report['status']='passed';report['startupRSSKiB']=int(subprocess.check_output(['ps','-p',str(owner.processes[-1].pid),'-o','rss='],text=True).strip())
        except BaseException:
            owner.log.flush();print((owner.directory/'packaged.log').read_text(errors='replace')[-7000:]);raise
        finally:
            cleanup_errors=[]
            try:
                for config in configs:
                    try:owner.api('/model-configs/'+config['id'],method='DELETE')
                    except Exception as error:cleanup_errors.append({'configurationId':config['id'],'errorType':type(error).__name__})
            finally:
                owner.close()
            if cleanup_errors:
                args.output.parent.mkdir(parents=True,exist_ok=True)
                args.output.with_suffix('.cleanup.json').write_text(json.dumps(cleanup_errors,indent=2)+'\n')
                raise AssertionError('Artificial configuration cleanup unconfirmed; exact IDs saved for remediation')
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--archive',type=Path,required=True);parser.add_argument('--builder-root',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);verify(parser.parse_args())
