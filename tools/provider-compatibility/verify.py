"""Explicitly authorized, artificial-data compatibility checks through app APIs.

Default command prints a receiver-pinned plan only. Keys stay in the application's
Keychain; this tool never loads credentials or reads agent authentication files.
"""
import argparse
import base64
from datetime import datetime,timezone
from decimal import Decimal,InvalidOperation
import hashlib
import importlib.metadata
import io
import ipaddress
import json
from pathlib import Path
import platform
import re
import subprocess
import threading
import time
from urllib.parse import urlsplit
import uuid

import httpx
from openpyxl import load_workbook

ARTIFICIAL_CSV=b'value\n1\n2\n6\n'
PROTOCOLS=['openai-responses','anthropic-messages']
ERROR_RESPONSE_MARKER='ARTIFICIAL_RAW_AUTH_RESPONSE_NOT_FOR_HISTORY'
INVALID_KEY='silicolab-invalid-compatibility-sentinel'
TERMINAL={'answered','cancelled','failed','paused','needs_decision','result_unknown'}
CASE_CHECKS={}
for protocol in PROTOCOLS:
    CASE_CHECKS['science-'+protocol]=('completed','receiver','streaming','toolAndArtifacts','resultReference','usage','artifactBytes')
    alternate=PROTOCOLS[1-PROTOCOLS.index(protocol)]
    CASE_CHECKS['switch-'+protocol+'-to-'+alternate]=('nextReceiver','completed','historyRetained','noRepeatedPython','resultReference','artifactBytes')
    CASE_CHECKS['cancel-'+protocol]=('observedInflightThenCancelled','sameReceiver','noTool')
    CASE_CHECKS['invalid-key-'+protocol]=('authenticationRejected','sameReceiver','noTool','noAutomaticRetry','safeError')
SCIENCE_CASE_ORDER=['science-openai-responses','switch-openai-responses-to-anthropic-messages','science-anthropic-messages','switch-anthropic-messages-to-openai-responses']
GOAL=('仅分析已固定的人工CSV，使用python_analysis一次。INPUTS是已提供的pathlib.Path，读取INPUTS/"input-001.csv"的value列；'
      '只用预装csv/json/math/statistics/openpyxl，禁止网络、安装、子进程。sum用math.fsum，mean用statistics.mean。'
      '在当前输出目录实际写summary.json，字段count,sum,mean；summary.xlsx仅一个工作表，第一行count,sum,mean，第二行三个数值。'
      '成功后不再运行工具，引用summary.json和summary.xlsx，用count=..., sum=..., mean=...解释实际结果。')
FOLLOWUP='请根据本对话上一项分析的真实工具结果，引用summary.json和summary.xlsx，明确写出count=..., sum=..., mean=...。不得重新调用工具或生成新产物。'
CANCEL_GOAL='这是一项人工取消验收。不要调用工具，逐行列出从1到2000的编号与对应英文拼写，持续输出，不提前总结。'


def sha256(data):return hashlib.sha256(data).hexdigest()
def now():return datetime.now(timezone.utc).isoformat()
def canonical(value):return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()


def case_passed(case):
    required=CASE_CHECKS.get(case.get('case'));checks=case.get('checks')
    return bool(required and isinstance(checks,dict) and set(checks)==set(required) and all(checks[key] is True for key in required))


def report_passed(report):
    cases=report.get('cases',[])
    return report.get('evidenceKind')=='real' and len(cases)==len(CASE_CHECKS) and {c.get('case') for c in cases}==set(CASE_CHECKS) and all(case_passed(case) for case in cases)


def build_plan(configurations,selected_ids,*,evidence_kind='real',allow_local_gateway=False,resume_prefix=None):
    if len(selected_ids)!=2 or len(set(selected_ids))!=2:raise ValueError('需要两个不同的应用配置ID')
    chosen=[]
    for identity in selected_ids:
        row=next((c for c in configurations if c['id']==identity),None)
        if not row or not row.get('hasKey'):raise ValueError('所选应用配置缺少key；请在SilicoLab设置中填写，不要发到聊天')
        if row.get('protocol') not in PROTOCOLS:raise ValueError('明确选择Responses和Messages两类协议')
        url=urlsplit(row['baseUrl'])
        loopback=url.hostname in ('localhost','127.0.0.1','::1')
        try:loopback=loopback or ipaddress.ip_address(url.hostname or '').is_loopback
        except ValueError:pass
        if url.username or url.password or url.query or url.fragment or not url.hostname:raise ValueError('端点不能包含认证或查询参数')
        if evidence_kind=='real' and ((loopback and not allow_local_gateway) or (url.scheme!='https' and not (loopback and allow_local_gateway and url.scheme=='http'))):raise ValueError('真实验收不能用回环假端点代替；真实本机网关需明确加入计划范围')
        if evidence_kind not in ('real','fixture'):raise ValueError('未知证据类型')
        chosen.append({k:row[k] for k in ('id','protocol','baseUrl','modelId')})
    chosen.sort(key=lambda c:PROTOCOLS.index(c['protocol']))
    if [c['protocol'] for c in chosen]!=PROTOCOLS:raise ValueError('必须各选一个Responses与Messages配置')
    plan={'format':'silicolab-provider-compatibility-v1','evidenceKind':evidence_kind,'allowLocalGateway':allow_local_gateway,'configurations':chosen,'protocols':PROTOCOLS,
          'applicationSourceSha256':source_identity(Path(__file__).resolve().parents[2])['sha256'],
          'runnerSha256':sha256(Path(__file__).read_bytes()),'promptSha256':sha256(canonical([GOAL,FOLLOWUP,CANCEL_GOAL])),
          'inputSha256':sha256(ARTIFICIAL_CSV),'inputDescription':'人工CSV value列1、2、6','tools':['python_analysis'],
          'cases':['each-protocol-science','same-conversation-switch-both-directions','each-protocol-cancel','each-endpoint-invalid-key'],
          'maximumTasks':8,'taskTimeoutSeconds':90,'requestObservationStop':20,
          'limits':'8个任务；每项最多90秒。观测到本轮20次请求则取消；观测限额不是原子HTTP计数上限。既有模型瞬态错误最多3次尝试。不购费、不调用科学远端、不读取私人文件。'}
    if allow_local_gateway:plan['limits']+=' 用户明确允许本机网关实际调用；记录的是网关协议与路由模型，不能推断上游供应商或模型身份，也不把fixture当真实。'
    if resume_prefix is not None:
        cases=resume_prefix.get('cases',[]);previous=resume_prefix.get('plan',{})
        if (previous.get('configurations')!=chosen or previous.get('evidenceKind')!=evidence_kind or previous.get('applicationSourceSha256')!=plan['applicationSourceSha256'] or
            len(cases) not in (2,4) or [c.get('case') for c in cases]!=SCIENCE_CASE_ORDER[:len(cases)] or
            not all(case_passed(c) and c.get('status')=='answered' for c in cases)):
            raise ValueError('只允许继续已完整核对的前两项或四项；不重放失败任务或变更接收方/应用源码')
        plan['resumePrefixSha256']=sha256(canonical(resume_prefix));plan['maximumTasks']=8-len(cases)
        plan['limits']=plan['limits'].replace('8个任务',f'仅剩余{plan["maximumTasks"]}个任务，原{len(cases)}项不重跑')
    plan['digest']=sha256(canonical(plan));return plan


class Application:
    def __init__(self,origin):
        parsed=urlsplit(origin)
        if parsed.scheme!='http' or parsed.hostname not in ('127.0.0.1','localhost','::1') or parsed.path not in ('','/'):raise ValueError('只连接指定SilicoLab本机应用')
        self.base=origin.rstrip('/')+'/api/workbench';self.client=httpx.Client(timeout=15,trust_env=False)
    def request(self,path,body=None,method=None):
        response=self.client.request(method or ('POST' if body is not None else 'GET'),self.base+path,json=body)
        if response.status_code>=400:raise ValueError('应用操作失败，HTTP '+str(response.status_code)+'；原记录保留')
        return response.json()
    def bytes(self,path):
        response=self.client.get(self.base+path);response.raise_for_status();return response.content
    def close(self):self.client.close()


class PartialEvidence:
    """Aggregate hashes only, scoped to one task and each in-flight request."""
    def __init__(self,task_id,*,sensitive=()):
        self.task_id=task_id;self.sensitive=sensitive;self.partials={};self.events=0;self.leak=False
    def add(self,state):
        self.events+=1
        self.leak=self.leak or any(value in canonical(state).decode() for value in self.sensitive)
        task=next((t for t in state.get('tasks',[]) if t['id']==self.task_id and t['status']=='running'),None)
        requests=[r for r in (task or {}).get('modelRequests',[]) if r['status']=='running']
        partial=state.get('partial',[])
        if len(requests)==1 and partial:
            self.partials.setdefault(requests[0]['id'],set()).add(sha256(canonical(partial)))
    def snapshot(self):
        rows=[{'taskId':self.task_id,'requestId':identity,'distinctPartials':len(parts)} for identity,parts in self.partials.items()]
        return {'requests':rows,'applicationEvents':self.events,'streaming':any(r['distinctPartials']>=2 for r in rows),'sensitiveContentObserved':self.leak}


class StreamObservation:
    """Observe actual application SSE partials, never infer streaming from final text."""
    def __init__(self,application,chat_url,task_id,*,sensitive=()):
        self.application=application;self.path=chat_url+'/events';self.stop=threading.Event();self.evidence=PartialEvidence(task_id,sensitive=sensitive);self.lock=threading.Lock()
        self.thread=threading.Thread(target=self._run,name='compatibility-sse-observer',daemon=True);self.thread.start()
    def _run(self):
        with httpx.Client(timeout=httpx.Timeout(5,read=1),trust_env=False) as client:
            while not self.stop.is_set():
                try:
                    with client.stream('GET',self.application.base+self.path) as response:
                        response.raise_for_status()
                        for line in response.iter_lines():
                            if self.stop.is_set():break
                            if not line.startswith('data: '):continue
                            state=json.loads(line[6:])
                            with self.lock:self.evidence.add(state)
                except (httpx.HTTPError,ValueError):self.stop.wait(.1)
    def snapshot(self):
        with self.lock:return self.evidence.snapshot()
    def close(self):self.stop.set();self.thread.join(timeout=7)


def assistant_text(state):
    parts=[]
    for entry in state.get('entries',[]):
        for message in entry.get('model',[]):
            if message.get('role')!='assistant':continue
            content=message.get('content',[])
            parts.append(content if isinstance(content,str) else ''.join(b.get('text','') for b in content if b.get('type')=='text'))
    return '\n'.join(parts)


def request_evidence(task,config):
    allowed=('id','configurationId','provider','protocol','endpoint','model','status','createdAt','finishedAt','attempt','parameters','usage','cost')
    requests=[{key:r[key] for key in allowed if key in r} for r in task.get('modelRequests',[])]
    same=bool(requests) and all(r.get('configurationId')==config['id'] and r.get('protocol')==config['protocol'] and r.get('endpoint')==config['baseUrl'] and r.get('model')==config['modelId'] for r in requests)
    return {'taskId':task['id'],'status':task['status'],'executionContract':task.get('executionContract'),'sameReceiver':same,'requests':requests}


def references_result(text):
    if not all(name in text for name in ('summary.json','summary.xlsx')):return False
    for field,number in [('count',3),('sum',9),('mean',3)]:
        tokens=re.findall(r'\b'+field+r'\s*[=:：]\s*([^\s,，;；。:：)）\]]+)',text)
        try:
            if not tokens or any(Decimal(token.rstrip('。*`'))!=number for token in tokens):return False
        except InvalidOperation:return False
    return True


def authentication_checks(task,state,config,stream,sentinel):
    evidence=request_evidence(task,config)
    return {'authenticationRejected':task['status']=='paused' and any('凭据无效' in r.get('error','') for r in task.get('modelRequests',[])),
            'sameReceiver':evidence['sameReceiver'],'noTool':task['toolExecutions']==[],'noAutomaticRetry':len(evidence['requests'])==1,
            'safeError':not stream.get('sensitiveContentObserved',True) and not any(value in canonical({'task':task,'state':state}).decode() for value in (sentinel,ERROR_RESPONSE_MARKER))}


def science_receipt(task,state,config,stream,*,evidence_kind):
    result=request_evidence(task,config);executions=task.get('toolExecutions',[])
    execution=executions[0] if len(executions)==1 else {};summary=execution.get('result',{})
    artifacts=execution.get('artifacts',[])
    checks={'completed':task['status']=='answered','receiver':result['sameReceiver'],'streaming':stream.get('streaming') is True,
            'toolAndArtifacts':len(executions)==1 and execution.get('tool')=='python_analysis' and execution.get('status')=='completed' and summary.get('checks',{}).get('independent_numeric_summary') is True and summary.get('summary')=={'count':3,'sum':9,'mean':3} and {'analysis.py','summary.json','summary.xlsx'}<=set(a['name'] for a in artifacts),
            'resultReference':references_result(assistant_text(state)),
            'usage':bool(result['requests']) and all(all(type(r.get('usage',{}).get(k)) in (int,float) and r['usage'][k]>=0 for k in ('input','output')) for r in result['requests'] if r['status']=='completed'),
            'artifactBytes':False}
    return {**result,'checks':checks,'stream':stream,'artifacts':[{k:a[k] for k in ('name','sha256','size') if k in a} for a in artifacts],
            'realCompatibilityPassed':False,'evidenceKind':evidence_kind}


def verify_artifacts(application,url,task):
    executions=task.get('toolExecutions',[])
    if len(executions)!=1:return False
    execution=executions[0];files={}
    for artifact in execution.get('artifacts',[]):
        name=artifact['name']
        if name not in ('analysis.py','summary.json','summary.xlsx'):continue
        data=application.bytes(url+'/tasks/'+task['id']+'/artifacts/'+execution['id']+'/'+name)
        if sha256(data)!=artifact['sha256']:return False
        files[name]=data
    if set(files)!={'analysis.py','summary.json','summary.xlsx'}:return False
    if json.loads(files['summary.json'])!={'count':3,'sum':9,'mean':3}:return False
    book=load_workbook(io.BytesIO(files['summary.xlsx']),read_only=True,data_only=True)
    try:return len(book.sheetnames)==1 and list(book.active.values)==[('count','sum','mean'),(3,9,3)] and bool(files['analysis.py'].strip())
    finally:book.close()


def run_suite(application,plan,workspace,approved_digest,*,resume_prefix=None):
    if not approved_digest or approved_digest!=plan['digest']:raise ValueError('未取得匹配计划的明确实际调用授权；未创建任务')
    expected=build_plan(application.request('/model-configs'),[c['id'] for c in plan['configurations']],evidence_kind=plan['evidenceKind'],allow_local_gateway=plan['allowLocalGateway'],resume_prefix=resume_prefix)
    if expected!=plan:raise ValueError('模型接收方或验收计划已变更，原授权不适用')
    projects=application.request('/projects');prefix_cases=[]
    if resume_prefix is not None:
        if len(projects)!=1:raise ValueError('继续验收必须只有原人工项目')
        project=application.request('/projects/'+projects[0]['id']);url='/projects/'+project['id']
        ids=[c['taskId'] for c in resume_prefix['cases']]
        if [t['id'] for t in project['tasks']]!=ids:raise ValueError('原项目任务清单已变更；未创建后续任务')
        for index,(case,task) in enumerate(zip(resume_prefix['cases'],project['tasks'])):
            config=plan['configurations'][[0,1,1,0][index]]
            evidence=request_evidence(task,config)
            if task['status']!='answered' or not evidence['sameReceiver'] or evidence['requests']!=case['requests'] or evidence['executionContract']!=case.get('executionContract'):raise ValueError('原请求来源或状态与收据不符')
            state=application.request(url+'/conversations/'+task['conversationId'])
            if not references_result(assistant_text({'entries':[e for e in state['entries'] if str(e['id']).startswith(task['id']+'-')]})):raise ValueError('原结果引用未通过重新核对')
        for offset in range(0,len(ids),2):
            first,second=project['tasks'][offset:offset+2]
            if first['conversationId']!=second['conversationId'] or second['toolExecutions'] or not verify_artifacts(application,url,first):raise ValueError('原工具、对话或产物不符')
            artifacts=[{k:a[k] for k in ('name','sha256','size') if k in a} for a in first['toolExecutions'][0]['artifacts']]
            if artifacts!=resume_prefix['cases'][offset]['artifacts']:raise ValueError('原产物来源变更')
        prefix_cases=json.loads(json.dumps(resume_prefix['cases']))
        for case in prefix_cases:case['recheckedBeforeContinuation']=now()
    elif projects:raise ValueError('验收只使用无既有项目的独立应用目录，不能附着私人项目')
    workspace=Path(workspace);workspace.mkdir(parents=True,exist_ok=False)
    if resume_prefix is None:
        project_path=workspace/'artificial-project';project_path.mkdir()
        project=application.request('/projects',{'path':str(project_path.resolve()),'name':'两协议人工兼容验收'})
        url='/projects/'+project['id']
    report={'format':plan['format'],'startedAt':now(),'plan':plan,'versions':versions(),'cases':prefix_cases,'evidenceKind':plan['evidenceKind'],'realCompatibilityPassed':False}
    out=workspace/'compatibility-report.json';active=None;pending_submission=None;observers=[];current_config=None
    def save():out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    def select(chat,config):
        nonlocal current_config
        application.request(url+'/conversations/'+chat+'/model',{'configurationId':config['id']},'PUT');current_config=config
    def start(chat,goal,**extra):
        nonlocal active,pending_submission
        if application.request(url)['modelRequests']>=plan['requestObservationStop']:raise ValueError('本轮请求达到观测停止线，未发起下一任务')
        pending_submission={'conversationId':chat,'requestId':str(uuid.uuid4())}
        report['pendingSubmission']=dict(pending_submission);save()
        task=application.request(url+'/tasks',{**pending_submission,'goal':goal,'acceptanceModel':{k:current_config[k] for k in ('id','protocol','baseUrl','modelId')},**extra})
        active=task['id'];pending_submission=None;report.pop('pendingSubmission',None);save()
        return task['id']
    def observe(chat,identity,*,cancel=False,sensitive=()):
        deadline=time.monotonic()+plan['taskTimeoutSeconds'];cancelled=False
        stream=StreamObservation(application,url+'/conversations/'+chat,identity,sensitive=sensitive);observers.append(stream)
        try:
            while time.monotonic()<deadline:
                state=application.request(url+'/conversations/'+chat);task=next(t for t in state['tasks'] if t['id']==identity)
                if state['modelRequests']>=plan['requestObservationStop']:raise ValueError('本轮请求达到观测停止线，停止后续验收')
                if task['status'] in TERMINAL:return task,state,stream.snapshot(),cancelled
                if cancel and not cancelled and any(r['status']=='running' for r in task.get('modelRequests',[])) and any(r['distinctPartials'] for r in stream.snapshot()['requests']):
                    application.request(url+'/tasks/'+identity+'/cancel',{},'POST');cancelled=True
                time.sleep(.04)
            raise TimeoutError('人工验收任务超过明确时限')
        finally:stream.close()
    try:
        configs=plan['configurations']
        for index,config in list(enumerate(configs))[len(prefix_cases)//2:]:
            chat=application.request(url+'/conversations',{'title':config['protocol']+'人工分析'})['id'];select(chat,config)
            source=application.request(url+'/files',{'name':'artificial.csv','contentBase64':base64.b64encode(ARTIFICIAL_CSV).decode()})
            fixed=application.request(url+'/input-versions',{'conversationId':chat,'fileIds':[source['id']]})
            scope={'workflow':'python_analysis','authorizedTools':['python_analysis'],'inputVersionId':fixed['id']}
            identity=start(chat,GOAL,**scope);task,state,stream,_=observe(chat,identity);active=None
            science=science_receipt(task,state,config,stream,evidence_kind=plan['evidenceKind']);science['case']='science-'+config['protocol']
            science['checks']['artifactBytes']=verify_artifacts(application,url,task)
            science['realCompatibilityPassed']=plan['evidenceKind']=='real' and case_passed(science);report['cases'].append(science);save()
            if task['status']!='answered' or task.get('modelBoundaryError'):raise ValueError('验收任务已暂停或未正常完成；保留结果，停止本轮并重新确认计划')
            if not science['checks']['toolAndArtifacts'] or not science['checks']['artifactBytes']:raise ValueError('真实科学工具或产物未通过；保留失败记录，不继续发送')
            completed=canonical(task['toolExecutions']);history={e['id']:sha256(canonical(e)) for e in state['entries']}
            alternate=configs[1-index];select(chat,alternate)
            follow=start(chat,FOLLOWUP,**scope);next_task,next_state,follow_stream,_=observe(chat,follow);active=None
            original=next(t for t in next_state['tasks'] if t['id']==identity)
            checks={'nextReceiver':request_evidence(next_task,alternate)['sameReceiver'],'completed':next_task['status']=='answered',
                    'historyRetained':all(any(e['id']==entry and sha256(canonical(e))==digest for e in next_state['entries']) for entry,digest in history.items()),
                    'noRepeatedPython':canonical(original['toolExecutions'])==completed and next_task['toolExecutions']==[],
                    'resultReference':references_result(assistant_text({'entries':[e for e in next_state['entries'] if e['id'] not in history]})),
                    'artifactBytes':verify_artifacts(application,url,original)}
            report['cases'].append({**request_evidence(next_task,alternate),'case':'switch-'+config['protocol']+'-to-'+alternate['protocol'],'checks':checks,'stream':follow_stream});save()
            if not all(checks.values()):raise ValueError('跨协议历史或工具复用验收未通过；停止后续发送')
        for config in configs:
            chat=application.request(url+'/conversations',{'title':config['protocol']+'人工取消'})['id'];select(chat,config)
            identity=start(chat,CANCEL_GOAL,authorizedTools=['synthetic_count']);task,state,stream,requested=observe(chat,identity,cancel=True);active=None
            evidence=request_evidence(task,config)
            checks={'observedInflightThenCancelled':requested and task['status']=='cancelled' and any(r['status']=='cancelled' for r in evidence['requests']),'sameReceiver':evidence['sameReceiver'],'noTool':task['toolExecutions']==[]}
            report['cases'].append({**evidence,'case':'cancel-'+config['protocol'],'checks':checks,'stream':stream});save()
        for config in configs:
            invalid=application.request('/model-configs',{'name':'人工无效key验收','protocol':config['protocol'],'baseUrl':config['baseUrl'],'modelId':config['modelId'],'apiKey':INVALID_KEY})
            try:
                chat=application.request(url+'/conversations',{'title':config['protocol']+'人工认证错误'})['id'];select(chat,invalid)
                identity=start(chat,'人工认证错误验收；不要调用工具，只答OK。',authorizedTools=['synthetic_count'])
                sentinel=INVALID_KEY
                task,state,stream,_=observe(chat,identity,sensitive=(sentinel,ERROR_RESPONSE_MARKER));active=None;evidence=request_evidence(task,invalid)
                checks=authentication_checks(task,state,invalid,stream,sentinel)
                report['cases'].append({**evidence,'case':'invalid-key-'+config['protocol'],'checks':checks});save()
            finally:application.request('/model-configs/'+invalid['id'],method='DELETE')
        report['realCompatibilityPassed']=report_passed(report)
    except (Exception,KeyboardInterrupt) as error:
        report['failureType']=type(error).__name__;report['stoppedWithoutRetry']=True
        if pending_submission and active is None:
            try:
                matches=[t for t in application.request(url)['tasks'] if t['requestId']==pending_submission['requestId'] and t['conversationId']==pending_submission['conversationId']]
                if len(matches)==1:active=matches[0]['id'];report['reconciledLostSubmission']=True
                else:report['submissionOutcomeUnconfirmed']=True
            except Exception:report['submissionOutcomeUnconfirmed']=True
        if active:
            try:application.request(url+'/tasks/'+active+'/cancel',{},'POST')
            except Exception:report['cancellationUnconfirmed']=True
            try:
                interrupted=next(t for t in application.request(url)['tasks'] if t['id']==active)
                report['cases'].append({**request_evidence(interrupted,current_config),'case':'interrupted-check','checks':{'completed':False}})
            except Exception:report['lastRequestEvidenceUnavailable']=True
        raise
    finally:
        for observer in observers:observer.close()
        report['finishedAt']=now();save()
    return report


def source_identity(root):
    root=Path(root);files={}
    for folder,suffixes in [('backend/src',('.py','.json','.xlsx')),('runtime/src',('.ts',)),('frontend/src',('.ts','.tsx','.css')),('frontend/dist',('.js','.css','.html'))]:
        for path in sorted((root/folder).rglob('*')):
            if path.is_file() and path.suffix in suffixes:files[path.relative_to(root).as_posix()]=sha256(path.read_bytes())
    for name in ('backend/pyproject.toml','runtime/package-lock.json','frontend/package-lock.json'):
        if (root/name).is_file():files[name]=sha256((root/name).read_bytes())
    return {'sha256':sha256(canonical(files)),'files':len(files)}


def versions():
    root=Path(__file__).resolve().parents[2]
    package=json.loads((root/'runtime/package-lock.json').read_text())
    return {'python':platform.python_version(),'macOS':platform.mac_ver()[0],'machine':platform.machine(),
            'backend':{name:importlib.metadata.version(name) for name in ('fastapi','uvicorn','httpx','openpyxl')},
            'runtime':{name:package['packages']['node_modules/'+name]['version'] for name in ('@earendil-works/pi-ai','@earendil-works/pi-durable')},
            'node':subprocess.check_output(['node','--version'],text=True).strip(),
            'applicationSources':source_identity(root),'runtimeLockSha256':sha256((root/'runtime/package-lock.json').read_bytes()),'runnerSha256':sha256(Path(__file__).read_bytes())}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--app-directory',type=Path,required=True,help='独立人工验收应用目录，不是私人项目目录')
    parser.add_argument('--responses-config',required=True);parser.add_argument('--messages-config',required=True)
    parser.add_argument('--resume-prefix',type=Path,help='已只读复核的前两项或四项收据；计划纳入摘要，执行再次核对，不重跑原模型或工具')
    parser.add_argument('--workspace',type=Path,help='新的人工输出目录；已存在则拒绝')
    parser.add_argument('--allow-local-gateway',action='store_true',help='仅在用户明确选择真实本机网关时纳入计划；不改变fixture证据类型')
    parser.add_argument('--approve-plan-sha256',help='仅在已取得用户明确实际调用授权后使用')
    args=parser.parse_args()
    from silicolab.desktop_lifecycle import service_info
    service=service_info(args.app_directory)
    if not service:raise ValueError('独立验收应用尚未启动；不自动启动私人后台')
    with httpx.Client(timeout=5,trust_env=False) as client:
        profile=client.get(service['origin']+'/api/desktop/status').json()
    if profile.get('isolatedResources') is not True:raise ValueError('需以--isolated-resources启动专用验收应用，不能发送本机已发现资源')
    application=Application(service['origin'])
    try:
        prefix=json.loads(args.resume_prefix.read_text()) if args.resume_prefix else None
        plan=build_plan(application.request('/model-configs'),[args.responses_config,args.messages_config],allow_local_gateway=args.allow_local_gateway,resume_prefix=prefix)
        if not args.approve_plan_sha256:print(json.dumps({'plan':plan,'versions':versions(),'executed':False},ensure_ascii=False,indent=2));return
        if not args.workspace:raise ValueError('实际验收需新的人工输出目录')
        report=run_suite(application,plan,args.workspace,args.approve_plan_sha256,resume_prefix=prefix)
        report['versions']=versions();(args.workspace/'compatibility-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
        print(json.dumps({'realCompatibilityPassed':report['realCompatibilityPassed'],'cases':len(report['cases'])}))
        if not report['realCompatibilityPassed']:raise SystemExit(1)
    finally:application.close()


if __name__=='__main__':main()
